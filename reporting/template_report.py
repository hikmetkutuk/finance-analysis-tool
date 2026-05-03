from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from pathlib import Path
from statistics import mean
from typing import Any, Iterable, Optional, Sequence

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.comments import Comment
from openpyxl.utils.cell import range_boundaries
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import TableColumn

from data.macro_config import load_macro_config
from data.ratio_profile_config import build_ratio_profile_map, load_ratio_profile_config, resolve_ratio_profile
from data.sector_profile_config import build_sector_profile_maps, load_sector_profile_config
from data.valuation_profile_config import build_valuation_profile_maps, load_valuation_profile_config, resolve_valuation_weight_map


DEFAULT_TEMPLATE_PATH = Path(__file__).resolve().parent.parent / "data" / "report_template.xlsx"
CODE_COLUMN = "Kod"
SECTOR_COLUMN = "Sektör"
INDEX_COLUMN = "Endeks"
VALUE_COLUMN = "Değer"
SECTION_COLUMN = "Bölüm"
TOPIC_COLUMN = "Konu"
WEIGHT_COLUMN = "Ağırlık"
TR_VALUE_COLUMN = "TR"
US_VALUE_COLUMN = "US"
INDEX_XU100 = "XU100"
INDEX_XUTUM = "XUTUM"
INDEX_US = "US"
INA_VALUE_COLUMN = "Profesyonel Değer"
SHEET_HISSE = "Hisse"
SHEET_PUAN = "Puan"
SHEET_ENDEKS = INDEX_COLUMN
SHEET_SEKTOR = SECTOR_COLUMN
SHEET_VARS = "Vars"
SHEET_RASYO = "Rasyo"
SHEET_NOTLAR = "Notlar"
SHEET_KALITE = "Kalite"
OUTPUT_SHEETS = (SHEET_HISSE, SHEET_PUAN, SHEET_ENDEKS, SHEET_SEKTOR, SHEET_VARS, SHEET_RASYO, SHEET_NOTLAR, SHEET_KALITE)
VALUE_NUMBER_FORMAT = "#,##0.00"
PERCENT_NUMBER_FORMAT = "0.00%"
GREEN_FONT_COLOR = "FF23BEA8"
RED_FONT_COLOR = "FFCC0000"
DEFAULT_FONT_COLOR = "FF666666"
RASYO_PASS_THRESHOLD = 0.60
MIN_MODEL_PRICE_MULTIPLE = 0.05
MAX_MODEL_PRICE_MULTIPLE = 5.0
MIN_PROFESSIONAL_MODELS = 3
MIN_PUBLISHABLE_MODELS = 2
MIN_PUBLISHABLE_CONFIDENCE = 0.35
MIN_PUBLISHABLE_DATA_COMPLETENESS = 0.45
HIGH_PE_WARNING = 50.0
EXTREME_PE_WARNING = 100.0
HIGH_EV_EBITDA_WARNING = 35.0
EXTREME_EV_EBITDA_WARNING = 60.0
MIN_BOND_ADJUSTMENT_FACTOR = 0.5
MAX_BOND_ADJUSTMENT_FACTOR = 1.5
BOND_REFERENCE_YIELDS = {
    "tr": 0.20,
    "us": 0.04,
}
MIN_PUBLISHABLE_ANCHOR_SCORE = 0.35
MAX_PRICE_ANCHOR_GAP = 0.75
MAX_ANALYST_ANCHOR_GAP = 0.60

PUAN_CODE = "1-Kod"
PUAN_NAME = "2-İsim"
PUAN_EBITDA = "3-FAVÖK"
PUAN_ROE = "4-Özs. Kar. (ROE) (%) Yıllık"
PUAN_EV_EBITDA = "5-FD/ FAVÖK"
PUAN_PCF = "6-F/NA"
PUAN_PE = "7-FK"
PUAN_EPS = "8-HBK"
PUAN_FORWARD_EPS = "8-İleri HBK"
PUAN_FORWARD_PE = "8-İleri FK"
PUAN_PB = "9-PD / DD"
PUAN_PS = "10-PD / NS"
PUAN_BETA = "11-Beta"
PUAN_ASSET_GROWTH = "12-Aktif Büyüme"
PUAN_NET_INCOME_GROWTH = "13-Net Kar Büyüme(Yıllık)"
PUAN_NET_INCOME = "14-NDK(Yıllık)"
PUAN_OPERATING_INCOME = "15-EFK(Yıllık)"
PUAN_EQUITY = "16-Özkaynaklar"
PUAN_DEBT_SOURCE = "17-Borç Kaynak"
PUAN_PAID_IN_CAPITAL = "18-HÖS"
PUAN_NET_WORKING_CAPITAL = "19-NİS"
PUAN_NET_DEBT_TO_EBITDA = "20-NB / FAVÖK (Yıllık) (%)"
PUAN_ANALYST_TARGET_MEAN = "21-Analist Hedef Ort."
PUAN_ANALYST_TARGET_MEDIAN = "22-Analist Hedef Medyan"
PUAN_ANALYST_COUNT = "23-Analist Sayısı"
PUAN_EARNINGS_GROWTH = "24-Kâr Büyüme Bekl."
PUAN_REVENUE_GROWTH = "25-Ciro Büyüme Bekl."
PUAN_FREE_CASH_FLOW = "26-Serbest Nakit Akımı"
PUAN_OPERATING_CASH_FLOW = "27-Operasyonel Nakit Akımı"
PUAN_REVENUE = "28-Gelirler"
PUAN_GROSS_MARGIN = "29-Brüt Kar Marjı (%)"
PUAN_TOTAL_ASSETS = "30-Toplam Varlıklar"
PUAN_TOTAL_DEBT = "31-Toplam Borç"

RASYO_STOCK = "Hisse"
RASYO_CURRENT_RATIO = "Cari Oran"
RASYO_QUICK_RATIO = "Likit Oran"
RASYO_CASH_RATIO = "Nakit Oran"
RASYO_ROE = "Özsermaye Karlılığı (ROE) (%) Yıllık"
RASYO_PE = "FK - Fiyat Kazanç (Dönem Sonu)"
RASYO_PB = "PD / DD"
RASYO_EBITDA_GROWTH = "FAVÖK Büyüme (%) (Yıllık)"
RASYO_PEG = "Peg Oranı"
RASYO_ROIC = "Roic"
RASYO_ASSET_TURNOVER = "Aktif Devir Hızı"
RASYO_PROFILE = "Profil"
RASYO_SCORE_100 = "Oran Skoru (0-100)"
RASYO_SCORE = "Oran Skoru (0-1)"
RASYO_COVERAGE = "Veri Kapsamı"
RASYO_MISSING_PENALTY = "Eksik Veri Cezası"
RASYO_LIQUIDITY_SCORE = "Likidite Skoru"
RASYO_PROFITABILITY_SCORE = "Karlılık Skoru"
RASYO_VALUATION_SCORE = "Değerleme Skoru"
RASYO_GROWTH_SCORE = "Büyüme Skoru"
RASYO_EFFICIENCY_SCORE = "Verimlilik Skoru"

HISSE_PRICE = "Fiyat"
HISSE_PE = "FK"
HISSE_PB = "PD/DD"
HISSE_NIS = "Nis"
HISSE_RATIO = "Rasyo"
HISSE_AVERAGE = "D. Ort"
HISSE_ANALYST_TARGET = "Analist Hedef"
HISSE_UPSIDE = "GP %"
HISSE_ANALYST_UPSIDE = "Analist GP %"
HISSE_CONFIDENCE = "Güven"
HISSE_FUNDAMENTAL_QUALITY = "Temel Kalite Skoru"
HISSE_MODEL_COUNT = "D. Ort Model Sayısı"
HISSE_STATUS = "Sonuç Durumu"
HISSE_STATUS_NOTE = "Yayın Notu"

QUALITY_CODE = CODE_COLUMN
QUALITY_MODEL = "Model"
QUALITY_RAW_VALUE = "Ham Değer"
QUALITY_FILTERED_VALUE = "Filtrelenmiş Değer"
QUALITY_STATUS = "Durum"
QUALITY_REASON = "Gerekçe"
QUALITY_WEIGHT_PROFILE = "Ağırlık Profili"
QUALITY_BASE_WEIGHT = "Baz Ağırlık"
QUALITY_ADJUSTED_WEIGHT = "Ayarlı Ağırlık"
QUALITY_PRICE = HISSE_PRICE
QUALITY_SECTOR = SECTOR_COLUMN
QUALITY_INDEX = INDEX_COLUMN
QUALITY_SIGNAL_SCORE = "Sinyal Kalite Skoru"
QUALITY_SIGNAL_LABEL = "Sinyal Güven Seviyesi"
QUALITY_WARNING_COUNT = "Model Kalite Uyarı Sayısı"
QUALITY_MODEL_WARNINGS = "Model Kalite Uyarıları"
QUALITY_DATA_COMPLETENESS = "Veri Tamlık Skoru"
QUALITY_FUNDAMENTAL_SCORE = "Temel Kalite Skoru"
QUALITY_MODEL_COUNT = "Geçerli Model Sayısı"
QUALITY_INCLUDED_IN_FAIR_VALUE = "D. Ort Dahil"
QUALITY_RESULT_STATUS = HISSE_STATUS
QUALITY_RESULT_NOTE = HISSE_STATUS_NOTE
STATUS_PUBLISHABLE = "Yayınlanabilir"
STATUS_REVIEW = "İnceleme Gerekli"
STATUS_UNPUBLISHABLE = "Yayınlanamaz"
QUALITY_STATUS_FILTERED = "Filtrelendi"
QUALITY_STATUS_UNWEIGHTED = "Ağırlıksız"
QUALITY_STATUS_TRIMMED = "Trimlendi"
QUALITY_STATUS_USED = "Kullanıldı"
RATIO_CATEGORY_LIQUIDITY = "Likidite"
RATIO_CATEGORY_PROFITABILITY = "Karlılık"
RATIO_CATEGORY_VALUATION = "Değerleme"
RATIO_CATEGORY_GROWTH = "Büyüme"
RATIO_CATEGORY_EFFICIENCY = "Verimlilik"

SECTOR_PB = RASYO_PB
SECTOR_PE = "F/K"
SECTOR_FD = "FD"
SECTOR_EV_EBITDA = "Firma Değeri/Favök"

PUAN_COLUMNS = (
    PUAN_CODE,
    PUAN_NAME,
    PUAN_EBITDA,
    PUAN_ROE,
    PUAN_EV_EBITDA,
    PUAN_PCF,
    PUAN_PE,
    PUAN_EPS,
    PUAN_FORWARD_EPS,
    PUAN_FORWARD_PE,
    PUAN_PB,
    PUAN_PS,
    PUAN_BETA,
    PUAN_ASSET_GROWTH,
    PUAN_NET_INCOME_GROWTH,
    PUAN_NET_INCOME,
    PUAN_OPERATING_INCOME,
    PUAN_EQUITY,
    PUAN_DEBT_SOURCE,
    PUAN_PAID_IN_CAPITAL,
    PUAN_NET_WORKING_CAPITAL,
    PUAN_NET_DEBT_TO_EBITDA,
    PUAN_ANALYST_TARGET_MEAN,
    PUAN_ANALYST_TARGET_MEDIAN,
    PUAN_ANALYST_COUNT,
    PUAN_EARNINGS_GROWTH,
    PUAN_REVENUE_GROWTH,
    PUAN_FREE_CASH_FLOW,
    PUAN_OPERATING_CASH_FLOW,
    PUAN_REVENUE,
    PUAN_GROSS_MARGIN,
    PUAN_TOTAL_ASSETS,
    PUAN_TOTAL_DEBT,
)

RASYO_COLUMNS = (
    RASYO_STOCK,
    RASYO_PROFILE,
    RASYO_CURRENT_RATIO,
    RASYO_QUICK_RATIO,
    RASYO_CASH_RATIO,
    RASYO_ROE,
    RASYO_PE,
    RASYO_PB,
    RASYO_EBITDA_GROWTH,
    RASYO_PEG,
    RASYO_ROIC,
    RASYO_ASSET_TURNOVER,
    RASYO_SCORE_100,
    RASYO_SCORE,
    RASYO_COVERAGE,
    RASYO_MISSING_PENALTY,
    RASYO_LIQUIDITY_SCORE,
    RASYO_PROFITABILITY_SCORE,
    RASYO_VALUATION_SCORE,
    RASYO_GROWTH_SCORE,
    RASYO_EFFICIENCY_SCORE,
)

HISSE_COLUMNS = (
    CODE_COLUMN,
    SECTOR_COLUMN,
    INDEX_COLUMN,
    HISSE_PRICE,
    HISSE_PE,
    HISSE_PB,
    HISSE_NIS,
    HISSE_RATIO,
    "D1",
    "D2",
    "D3",
    "D4",
    "D5",
    "D6",
    "D7",
    "D8",
    "D9",
    "D10",
    "D11",
    "D12",
    HISSE_AVERAGE,
    HISSE_UPSIDE,
    HISSE_ANALYST_TARGET,
    HISSE_ANALYST_UPSIDE,
    HISSE_CONFIDENCE,
    HISSE_FUNDAMENTAL_QUALITY,
    HISSE_MODEL_COUNT,
    HISSE_STATUS,
    HISSE_STATUS_NOTE,
)

QUALITY_COLUMNS = (
    QUALITY_CODE,
    QUALITY_MODEL,
    QUALITY_RAW_VALUE,
    QUALITY_FILTERED_VALUE,
    QUALITY_STATUS,
    QUALITY_REASON,
    QUALITY_WEIGHT_PROFILE,
    QUALITY_BASE_WEIGHT,
    QUALITY_ADJUSTED_WEIGHT,
    QUALITY_INCLUDED_IN_FAIR_VALUE,
    QUALITY_PRICE,
    QUALITY_SECTOR,
    QUALITY_INDEX,
    QUALITY_SIGNAL_SCORE,
    QUALITY_SIGNAL_LABEL,
    QUALITY_WARNING_COUNT,
    QUALITY_MODEL_WARNINGS,
    QUALITY_DATA_COMPLETENESS,
    QUALITY_FUNDAMENTAL_SCORE,
    QUALITY_MODEL_COUNT,
    QUALITY_RESULT_STATUS,
    QUALITY_RESULT_NOTE,
)

HISSE_HEADER_COMMENTS = {
    "D1": "S(F/K) * HBK",
    "D2": "EFK * 10 / HOS",
    "D3": "NDK * 10 / HOS",
    "D4": "(Ozkaynaklar / HOS) * S(PD/DD)",
    "D5": "Ozkaynaklar / HOS",
    "D6": "İleri HBK ve İleri FK varsa İleri HBK * İleri FK; yoksa HBK * (1 + Aktif Buyume / 100) * FK",
    "D7": "((FAVOK * S(FD/FAVOK)) - Net Borc) / HOS",
    "D8": "S(F/K) * HBK, 2Y tahvile gore normalize edilmis faiz rejimi ayarli carpan modeli",
    "D9": "(EFK * (1 + terminal_growth) / (AOSM - terminal_growth)) / HOS",
    "D10": "(NDK * (1 + terminal_growth) / (AOSM - terminal_growth)) / HOS",
    "D11": "DCF/INA profesyonel deger. Pozitif olmayan degerler bos birakilir.",
    "D12": "(Ozkaynaklar / HOS) * ((ROE - terminal_growth) / (Ozkaynak Maliyeti - terminal_growth))",
}

D_FIELDS = tuple(f"D{index}" for index in range(1, 13))
WEIGHT_PROFILE_COLUMN = "Ağırlık Profili"
WEIGHT_NOTE_COLUMN = "Not"
WEIGHT_PROFILE_DEFAULT = "Genel Sanayi"
WEIGHT_PROFILE_FINANCIAL = "Finansal"
WEIGHT_PROFILE_GROWTH = "Büyüme/Teknoloji"
WEIGHT_PROFILE_REAL_ESTATE = "Gayrimenkul"
WEIGHT_PROFILE_HOLDING = "Holding"
WEIGHT_PROFILE_ENERGY_UTILITY = "Enerji Utility/Altyapı"
WEIGHT_PROFILE_ENERGY_EQUIPMENT = "Enerji Ekipman/Taahhüt"

SECTOR_PROFILE_CONFIG = load_sector_profile_config()
SECTOR_PROFILE_LOOKUP, SECTOR_PROFILE_KEYWORDS = build_sector_profile_maps(SECTOR_PROFILE_CONFIG)
VALUATION_PROFILE_CONFIG = load_valuation_profile_config()
MODEL_WEIGHTS, PROFILE_NOTES = build_valuation_profile_maps(VALUATION_PROFILE_CONFIG)
RATIO_PROFILE_CONFIG = load_ratio_profile_config()
RATIO_PROFILES = build_ratio_profile_map(RATIO_PROFILE_CONFIG)


@dataclass(frozen=True)
class MacroRates:
    market_key: str
    two_year_bond: Optional[float]
    market_premium: Optional[float]
    terminal_growth: Optional[float]
    cost_of_debt: Optional[float]
    tax_rate: Optional[float]
    risk_free_rate: Optional[float]


@dataclass(frozen=True)
class MacroRateBook:
    tr: MacroRates
    us: MacroRates

    def for_ticker(self, ticker: str) -> MacroRates:
        return self.tr if ticker.upper().endswith(".IS") else self.us


@dataclass(frozen=True)
class RowLookups:
    valuation: pd.DataFrame
    puan: pd.DataFrame
    rasyo: pd.DataFrame
    endeks: pd.DataFrame
    sektor: pd.DataFrame
    ina: pd.DataFrame


@dataclass(frozen=True)
class HisseInputs:
    code: str
    sector_name: str
    index_name: str
    market_key: str
    price: Optional[float]
    company_pe: Optional[float]
    company_pb: Optional[float]
    nis_sign: str
    ratio_score: Optional[float]
    sector_pe: Optional[float]
    sector_pb: Optional[float]
    sector_ev_ebitda: Optional[float]
    eps: Optional[float]
    forward_eps: Optional[float]
    forward_pe: Optional[float]
    roe: Optional[float]
    beta: Optional[float]
    debt_ratio: Optional[float]
    net_income_growth: Optional[float]
    earnings_growth: Optional[float]
    revenue_growth: Optional[float]
    net_income: Optional[float]
    operating_income: Optional[float]
    equity: Optional[float]
    paid_in_capital: Optional[float]
    ebitda: Optional[float]
    free_cash_flow: Optional[float]
    operating_cash_flow: Optional[float]
    net_debt: float
    asset_growth: Optional[float]
    dcf_value: Optional[float]
    analyst_target: Optional[float]
    analyst_count: Optional[float]
    revenue: Optional[float] = None
    gross_margin: Optional[float] = None
    total_assets: Optional[float] = None
    total_debt: Optional[float] = None


@dataclass(frozen=True)
class ValuationSummary:
    fair_value: Optional[float]
    confidence: Optional[float]
    model_count: int
    publishable: bool
    status: str
    note: str
    data_completeness: Optional[float]


@dataclass(frozen=True)
class RatioScoreSummary:
    profile_name: str
    normalized_score: Optional[float]
    score_100: Optional[float]
    coverage: Optional[float]
    missing_penalty: Optional[float]
    category_scores: dict[str, Optional[float]]


@dataclass(frozen=True)
class QualityRowContext:
    signal_score: Optional[float]
    signal_label: str
    warning_count: Optional[float]
    model_warnings: str
    data_completeness: float
    summary: ValuationSummary

SHEET_UPDATES_Q2_2026_XU100_ADD = {"CVKMD", "EUREN", "PAHOL", "PSGYO", "SARKY"}
SHEET_UPDATES_Q2_2026_XU100_REMOVE = {"EGEEN", "KCAER", "LINK", "TTRAK", "YEOTK"}


def _round_or_none(value: Any, digits: int = 2) -> Optional[float]:
    try:
        if value is None or pd.isna(value):
            return None
        return round(float(value), digits)
    except (TypeError, ValueError):
        return None


def _safe_float(value: Any) -> Optional[float]:
    try:
        if value is None or pd.isna(value):
            return None
        return float(value)
    except (TypeError, ValueError):
        return None


def _positive_or_none(value: Any) -> Optional[float]:
    number = _safe_float(value)
    if number is None or number <= 0:
        return None
    return number


def _normalized_code(value: Any) -> str:
    if value is None or pd.isna(value):
        return ""
    code = str(value).strip().upper()
    return code[:-3] if code.endswith(".IS") else code


def _safe_divide(numerator: Optional[float], denominator: Optional[float]) -> Optional[float]:
    if numerator is None or denominator in (None, 0):
        return None
    try:
        return numerator / denominator
    except (TypeError, ZeroDivisionError):
        return None


def _rate_decimal(value: Any) -> Optional[float]:
    rate = _safe_float(value)
    if rate is None:
        return None
    return rate / 100.0 if abs(rate) > 1 else rate


def _rate_percent_points(value: Any) -> Optional[float]:
    rate = _safe_float(value)
    if rate is None:
        return None
    return rate * 100.0 if abs(rate) <= 1 else rate


def _bond_adjusted_multiple_value(
    base_value: Any,
    two_year_bond: Optional[float],
    market_key: str,
) -> Optional[float]:
    value = _safe_float(base_value)
    rate = _rate_decimal(two_year_bond)
    reference_rate = _rate_decimal(BOND_REFERENCE_YIELDS.get(market_key))
    if value is None or value <= 0 or rate is None or rate <= 0 or reference_rate is None or reference_rate <= 0:
        return None
    adjustment_factor = reference_rate / rate
    adjustment_factor = max(MIN_BOND_ADJUSTMENT_FACTOR, min(MAX_BOND_ADJUSTMENT_FACTOR, adjustment_factor))
    return value * adjustment_factor


def _sanity_checked_model_value(value: Any, price: Optional[float]) -> Optional[float]:
    number = _safe_float(value)
    if number is None or number <= 0:
        return None
    if price is None or price <= 0:
        return number
    if number < price * MIN_MODEL_PRICE_MULTIPLE:
        return None
    if number > price * MAX_MODEL_PRICE_MULTIPLE:
        return None
    return number


def sanity_checked_valuation_points(values: dict[str, Optional[float]], price: Optional[float]) -> dict[str, Optional[float]]:
    return {model_key: _sanity_checked_model_value(value, price) for model_key, value in values.items()}


def _model_value_reason(raw_value: Any, filtered_value: Any, price: Optional[float]) -> str:
    raw_number = _safe_float(raw_value)
    filtered_number = _safe_float(filtered_value)
    if filtered_number is not None:
        return "kullanilabilir"
    if raw_number is None:
        return "girdi_eksik"
    if raw_number <= 0:
        return "pozitif_degil"
    if price is None or price <= 0:
        return "fiyat_eksik"
    if raw_number < price * MIN_MODEL_PRICE_MULTIPLE:
        return "fiyata_gore_cok_dusuk"
    if raw_number > price * MAX_MODEL_PRICE_MULTIPLE:
        return "fiyata_gore_cok_yuksek"
    return "filtre_disinda"


def _debt_weight_from_debt_source(value: Any) -> Optional[float]:
    debt_source = _rate_decimal(value)
    if debt_source is None:
        return None
    denominator = 1.0 + debt_source
    if denominator == 0:
        return None
    return max(0.0, min(1.0, debt_source / denominator))


def _discount_rate(
    debt_ratio: Optional[float],
    beta: Optional[float],
    cost_of_debt: Optional[float],
    tax_rate: Optional[float],
    risk_free_rate: Optional[float],
    market_premium: Optional[float],
) -> Optional[float]:
    if cost_of_debt is None or tax_rate is None or risk_free_rate is None or market_premium is None:
        return None
    normalized_debt_ratio = debt_ratio if debt_ratio is not None else 0.0
    equity_ratio = 1.0 - normalized_debt_ratio
    normalized_beta = beta if beta is not None else 1.0
    cost_of_equity = risk_free_rate + (normalized_beta * market_premium)
    after_tax_debt_cost = cost_of_debt * (1.0 - tax_rate)
    return (after_tax_debt_cost * normalized_debt_ratio) + (cost_of_equity * equity_ratio)


def _cost_of_equity(
    beta: Optional[float],
    risk_free_rate: Optional[float],
    market_premium: Optional[float],
) -> Optional[float]:
    if risk_free_rate is None or market_premium is None:
        return None
    normalized_beta = beta if beta is not None else 1.0
    return risk_free_rate + (normalized_beta * market_premium)


def _roe_justified_pb_value(
    book_value_per_share: Optional[float],
    roe: Optional[float],
    cost_of_equity: Optional[float],
    terminal_growth: Optional[float],
) -> Optional[float]:
    if (
        book_value_per_share is None
        or roe is None
        or cost_of_equity is None
        or terminal_growth is None
        or cost_of_equity <= terminal_growth
    ):
        return None
    justified_pb = (roe - terminal_growth) / (cost_of_equity - terminal_growth)
    if justified_pb <= 0:
        return None
    return book_value_per_share * justified_pb


def _mean_nonempty(values: Iterable[Any]) -> Optional[float]:
    numeric_values = [number for value in values if (number := _safe_float(value)) is not None]
    if not numeric_values:
        return None
    return mean(numeric_values)


def _weight_profile_for_sector(sector_name: str) -> str:
    normalized = str(sector_name or "").strip()
    if not normalized:
        return WEIGHT_PROFILE_DEFAULT
    direct = SECTOR_PROFILE_LOOKUP.get(normalized)
    if direct is not None:
        return direct
    lowered = normalized.casefold()
    for keyword, profile_name in SECTOR_PROFILE_KEYWORDS:
        if keyword in lowered:
            return profile_name
    return WEIGHT_PROFILE_DEFAULT


def _ratio_weight_factor(inputs: HisseInputs) -> float:
    if inputs.ratio_score is not None:
        if inputs.ratio_score < 0.30:
            return 0.55
        if inputs.ratio_score < 0.50:
            return 0.75
    return 1.0


def _pe_weight_factor(model_key: str, inputs: HisseInputs) -> float:
    if model_key not in {"D1", "D6"}:
        return 1.0

    pe_reference = inputs.forward_pe if inputs.forward_pe is not None else inputs.company_pe
    if pe_reference is None:
        return 1.0
    if pe_reference > EXTREME_PE_WARNING:
        return 0.25
    if pe_reference > HIGH_PE_WARNING:
        return 0.55
    return 1.0


def _trailing_pe_fallback_factor(model_key: str, inputs: HisseInputs) -> float:
    if model_key == "D6" and inputs.forward_pe is None and inputs.company_pe is not None and inputs.company_pe > HIGH_PE_WARNING:
        return 0.60
    return 1.0


def _ev_ebitda_weight_factor(model_key: str, inputs: HisseInputs) -> float:
    if model_key != "D7" or inputs.sector_ev_ebitda is None:
        return 1.0
    if inputs.sector_ev_ebitda > EXTREME_EV_EBITDA_WARNING:
        return 0.30
    if inputs.sector_ev_ebitda > HIGH_EV_EBITDA_WARNING:
        return 0.60
    return 1.0


def _growth_weight_factor(model_key: str, inputs: HisseInputs) -> float:
    growth_reference = inputs.earnings_growth if inputs.earnings_growth is not None else inputs.net_income_growth
    if model_key in {"D2", "D3", "D6", "D9", "D10"} and growth_reference is not None and growth_reference < 0:
        return 0.60
    return 1.0


def _debt_weight_factor(model_key: str, inputs: HisseInputs) -> float:
    if model_key in {"D9", "D10", "D12"} and inputs.debt_ratio is not None and inputs.debt_ratio > 0.60:
        return 0.75
    return 1.0


def _model_weight_factor(model_key: str, inputs: HisseInputs) -> float:
    factors = (
        _ratio_weight_factor(inputs),
        _pe_weight_factor(model_key, inputs),
        _trailing_pe_fallback_factor(model_key, inputs),
        _ev_ebitda_weight_factor(model_key, inputs),
        _growth_weight_factor(model_key, inputs),
        _debt_weight_factor(model_key, inputs),
    )
    factor = 1.0
    for current_factor in factors:
        factor *= current_factor
    return factor


def _adjusted_model_weights(values: dict[str, Optional[float]], inputs: HisseInputs, weight_profile: str) -> dict[str, float]:
    base_weights = resolve_valuation_weight_map(VALUATION_PROFILE_CONFIG, weight_profile, inputs.market_key)
    adjusted: dict[str, float] = {}
    for model_key, base_weight in base_weights.items():
        value = _safe_float(values.get(model_key))
        if value is None or value <= 0:
            continue
        adjusted[model_key] = base_weight * _model_weight_factor(model_key, inputs)
    return adjusted


def _active_model_values(values: dict[str, Optional[float]], model_keys: Iterable[str]) -> dict[str, float]:
    active_values: dict[str, float] = {}
    for model_key in model_keys:
        value = _safe_float(values.get(model_key))
        if value is not None:
            active_values[model_key] = value
    return active_values


def _weighted_average(values: dict[str, Optional[float]], weights: dict[str, float]) -> Optional[float]:
    weighted_sum = 0.0
    active_weight = 0.0
    for model_key, weight in weights.items():
        value = _safe_float(values.get(model_key))
        if value is None or value <= 0 or weight <= 0:
            continue
        weighted_sum += value * weight
        active_weight += weight
    if active_weight <= 0:
        return None
    return weighted_sum / active_weight


def _model_dispersion(values: dict[str, Optional[float]], weights: dict[str, float], fair_value: Optional[float]) -> Optional[float]:
    if fair_value is None or fair_value <= 0 or not weights:
        return None
    active_values = list(_active_model_values(values, weights).values())
    if len(active_values) < 2:
        return None
    avg_abs_deviation = mean(abs(value - fair_value) for value in active_values)
    return avg_abs_deviation / fair_value


def _confidence_from_dispersion(dispersion: Optional[float]) -> float:
    if dispersion is None:
        return 0.55
    if dispersion <= 0.25:
        return 1.0
    if dispersion >= 1.25:
        return 0.10
    return max(0.10, 1.0 - ((dispersion - 0.25) / 1.0) * 0.90)


def _data_completeness_score(inputs: HisseInputs) -> float:
    fields = (
        inputs.price,
        inputs.paid_in_capital,
        inputs.eps,
        inputs.forward_eps,
        inputs.forward_pe,
        inputs.ebitda,
        inputs.net_income,
        inputs.operating_income,
        inputs.equity,
        inputs.sector_pe,
        inputs.sector_ev_ebitda,
        inputs.dcf_value,
        inputs.analyst_target,
        inputs.revenue,
        inputs.gross_margin,
        inputs.total_assets,
        inputs.total_debt,
    )
    return sum(1 for value in fields if _safe_float(value) is not None) / len(fields)


def _linear_min_score(value: Optional[float], soft_min: float, target_min: float) -> Optional[float]:
    if value is None:
        return None
    if value <= soft_min:
        return 0.10
    if value >= target_min:
        return 1.0
    return max(0.10, min(1.0, 0.10 + ((value - soft_min) / (target_min - soft_min)) * 0.90))


def _linear_max_score(value: Optional[float], target_max: float, soft_max: float) -> Optional[float]:
    if value is None:
        return None
    if value <= target_max:
        return 1.0
    if value >= soft_max:
        return 0.10
    return max(0.10, min(1.0, 1.0 - ((value - target_max) / (soft_max - target_max)) * 0.90))


def _fundamental_quality_score(inputs: HisseInputs) -> float:
    scores: list[tuple[float, float]] = []
    gross_margin_score = _linear_min_score(_safe_float(inputs.gross_margin), 10.0, 30.0)
    if gross_margin_score is not None:
        scores.append((gross_margin_score, 0.25))

    debt_burden = _safe_divide(inputs.total_debt, inputs.total_assets)
    debt_burden_score = _linear_max_score(debt_burden, 0.30, 0.80)
    if debt_burden_score is not None:
        scores.append((debt_burden_score, 0.20))

    asset_productivity = _safe_divide(inputs.revenue, inputs.total_assets)
    asset_productivity_score = _linear_min_score(asset_productivity, 0.20, 0.80)
    if asset_productivity_score is not None:
        scores.append((asset_productivity_score, 0.20))

    if inputs.ratio_score is not None:
        ratio_quality_score = 1.0 if inputs.ratio_score >= 0.60 else max(0.10, inputs.ratio_score / 0.60)
        scores.append((ratio_quality_score, 0.20))

    nis_quality_score = 0.50
    if inputs.nis_sign == "+":
        nis_quality_score = 1.0
    elif inputs.nis_sign == "-":
        nis_quality_score = 0.20
    scores.append((nis_quality_score, 0.15))

    if not scores:
        return 0.50
    total_weight = sum(weight for _, weight in scores)
    return sum(score * weight for score, weight in scores) / total_weight


def _relative_gap(reference_value: Optional[float], compared_value: Optional[float]) -> Optional[float]:
    reference = _safe_float(reference_value)
    compared = _safe_float(compared_value)
    if reference is None or compared is None or reference <= 0 or compared <= 0:
        return None
    return abs(compared - reference) / reference


def _anchor_gap_score(gap: Optional[float], max_gap: float) -> Optional[float]:
    if gap is None:
        return None
    if gap <= 0.15:
        return 1.0
    if gap >= max_gap:
        return 0.10
    return max(0.10, 1.0 - ((gap - 0.15) / (max_gap - 0.15)) * 0.90)


def _anchor_consistency_score(
    fair_value: Optional[float],
    price: Optional[float],
    analyst_target: Optional[float],
) -> float:
    scores: list[tuple[float, float]] = []
    price_score = _anchor_gap_score(_relative_gap(price, fair_value), MAX_PRICE_ANCHOR_GAP)
    analyst_score = _anchor_gap_score(_relative_gap(analyst_target, fair_value), MAX_ANALYST_ANCHOR_GAP)
    if price_score is not None:
        scores.append((price_score, 0.60))
    if analyst_score is not None:
        scores.append((analyst_score, 0.40))
    if not scores:
        return 0.50
    total_weight = sum(weight for _, weight in scores)
    return sum(score * weight for score, weight in scores) / total_weight


def _publication_decision(
    fair_value: Optional[float],
    confidence: Optional[float],
    model_count: int,
    data_score: float,
    anchor_score: float,
) -> tuple[bool, str, str]:
    if fair_value is None or model_count == 0:
        return False, STATUS_UNPUBLISHABLE, "gecerli_model_yok"
    if model_count < MIN_PUBLISHABLE_MODELS:
        return False, STATUS_REVIEW, "model_sayisi_yetersiz"
    if data_score < MIN_PUBLISHABLE_DATA_COMPLETENESS:
        return False, STATUS_REVIEW, "veri_tamligi_dusuk"
    if anchor_score < MIN_PUBLISHABLE_ANCHOR_SCORE:
        return False, STATUS_REVIEW, "piyasa_konsensus_uyumsuzlugu"
    if confidence is None or confidence < MIN_PUBLISHABLE_CONFIDENCE:
        return False, STATUS_REVIEW, "guven_dusuk"
    return True, STATUS_PUBLISHABLE, "yeterli_kanit"


def valuation_summary(values: dict[str, Optional[float]], inputs: HisseInputs, weight_profile: str) -> ValuationSummary:
    adjusted_weights = _adjusted_model_weights(values, inputs, weight_profile)
    active_weights = {
        model_key: weight
        for model_key, weight in adjusted_weights.items()
        if weight > 0
    }
    raw_fair_value = _weighted_average(values, active_weights)
    model_count = len(active_weights)
    data_score = _data_completeness_score(inputs)
    if raw_fair_value is None or model_count == 0:
        return ValuationSummary(None, None, 0, False, "Yayınlanamaz", "gecerli_model_yok", data_score)

    model_score = min(1.0, model_count / MIN_PROFESSIONAL_MODELS)
    ratio_score = inputs.ratio_score if inputs.ratio_score is not None else 0.50
    dispersion_score = _confidence_from_dispersion(_model_dispersion(values, active_weights, raw_fair_value))
    anchor_score = _anchor_consistency_score(raw_fair_value, inputs.price, inputs.analyst_target)
    fundamental_score = _fundamental_quality_score(inputs)
    confidence = (
        (model_score * 0.22)
        + (dispersion_score * 0.22)
        + (data_score * 0.16)
        + (ratio_score * 0.12)
        + (anchor_score * 0.14)
        + (fundamental_score * 0.14)
    )
    if model_count < 2:
        confidence *= 0.65
    normalized_confidence = max(0.0, min(1.0, confidence))
    preferred_quality_gate = (inputs.ratio_score is None or inputs.ratio_score >= 0.60) and inputs.nis_sign == "+"
    if not preferred_quality_gate:
        normalized_confidence = min(normalized_confidence, 0.35 + (fundamental_score * 0.45))
    if anchor_score < MIN_PUBLISHABLE_ANCHOR_SCORE:
        normalized_confidence = min(normalized_confidence, max(0.20, anchor_score))
    publishable, status, note = _publication_decision(
        raw_fair_value,
        normalized_confidence,
        model_count,
        data_score,
        anchor_score,
    )
    published_fair_value = raw_fair_value if publishable else None
    return ValuationSummary(published_fair_value, normalized_confidence, model_count, publishable, status, note, data_score)


def _median_or_none(series: pd.Series) -> Optional[float]:
    numeric = pd.to_numeric(series, errors="coerce").dropna()
    if numeric.empty:
        return None
    return _round_or_none(numeric.median())


def _column_index(columns: Sequence[str]) -> pd.Index:
    return pd.Index(columns)


def _clear_sheet(ws, start_row: int, start_col: int, end_col: Optional[int] = None) -> None:
    last_col = end_col if end_col is not None else ws.max_column
    for row in ws.iter_rows(min_row=start_row, max_row=ws.max_row, min_col=start_col, max_col=last_col):
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            cell.value = None


def _unmerge_sheet(ws) -> None:
    for merged_range in tuple(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged_range))


def _write_table(ws, frame: pd.DataFrame, header_row: int = 1, first_col: int = 1, include_header: bool = True) -> None:
    row_index = header_row
    if include_header:
        for col_offset, column_name in enumerate(frame.columns, start=first_col):
            ws.cell(row=row_index, column=col_offset).value = column_name
        row_index += 1
    for _, row in frame.iterrows():
        for col_offset, column_name in enumerate(frame.columns, start=first_col):
            value = row[column_name]
            ws.cell(row=row_index, column=col_offset).value = None if pd.isna(value) else value
        row_index += 1


def _sync_excel_table(ws, table_name: str, columns: Sequence[str], header_row: int, row_count: int) -> None:
    if table_name not in ws.tables:
        return
    table = ws.tables[table_name]
    end_col = get_column_letter(len(columns))
    end_row = max(header_row + row_count, header_row + 1)
    table.ref = f"A{header_row}:{end_col}{end_row}"
    table.tableColumns = [
        TableColumn(id=index, name=str(column_name))
        for index, column_name in enumerate(columns, start=1)
    ]
    if table.autoFilter is not None:
        table.autoFilter.ref = table.ref


def _copy_cell_style(source, target) -> None:
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.protection = copy(source.protection)
    target.number_format = source.number_format


def _align_hisse_column_styles(ws, columns: Sequence[str], header_row: int, row_count: int) -> None:
    headers = {column_name: index for index, column_name in enumerate(columns, start=1)}
    first_data_row = header_row + 1
    last_data_row = header_row + row_count
    data_style_column = headers.get("D1", 1)

    for column_name in (HISSE_RATIO, HISSE_ANALYST_TARGET, HISSE_UPSIDE, HISSE_ANALYST_UPSIDE, HISSE_CONFIDENCE, HISSE_MODEL_COUNT):
        column_index = headers.get(column_name)
        if column_index is None or column_index <= 1:
            continue
        _copy_cell_style(ws.cell(row=header_row, column=column_index - 1), ws.cell(row=header_row, column=column_index))
        for row_index in range(first_data_row, last_data_row + 1):
            _copy_cell_style(ws.cell(row=row_index, column=data_style_column), ws.cell(row=row_index, column=column_index))


def _apply_number_format_to_columns(
    ws,
    headers: dict[str, int],
    column_names: Sequence[str],
    number_format: str,
    first_row: int,
    last_row: int,
) -> None:
    for column_name in column_names:
        column_index = headers.get(column_name)
        if column_index is None:
            continue
        for row_index in range(first_row, last_row + 1):
            ws.cell(row=row_index, column=column_index).number_format = number_format


def _existing_column_indexes(headers: dict[str, int], column_names: Iterable[str]) -> set[int]:
    return {
        column_index
        for column_name in column_names
        if (column_index := headers.get(column_name)) is not None
    }


def _apply_header_comments(ws, columns: Sequence[str], comments: dict[str, str], header_row: int) -> None:
    for column_index, column_name in enumerate(columns, start=1):
        comment_text = comments.get(column_name)
        if comment_text is None:
            continue
        ws.cell(row=header_row, column=column_index).comment = Comment(comment_text, "finance-analysis-tool")


def _remove_conditional_formatting_for_columns(ws, column_indexes: set[int]) -> None:
    retained_rules = []
    cf_rules = getattr(ws.conditional_formatting, "_cf_rules")
    for conditional_range, rules in cf_rules.items():
        keep_range = True
        for cell_range in conditional_range.sqref.ranges:
            min_col, _, max_col, _ = range_boundaries(str(cell_range))
            if any(min_col <= column_index <= max_col for column_index in column_indexes):
                keep_range = False
                break
        if keep_range:
            retained_rules.append((conditional_range, rules))

    cf_rules.clear()
    for conditional_range, rules in retained_rules:
        cf_rules[conditional_range] = rules


def _apply_hisse_number_formats(ws, columns: Sequence[str], header_row: int, row_count: int) -> None:
    headers = {column_name: index for index, column_name in enumerate(columns, start=1)}
    first_data_row = header_row + 1
    last_data_row = header_row + row_count

    value_columns = tuple(f"D{index}" for index in range(1, 13)) + (HISSE_AVERAGE, HISSE_ANALYST_TARGET)
    percent_columns = (HISSE_UPSIDE, HISSE_ANALYST_UPSIDE, HISSE_CONFIDENCE, HISSE_FUNDAMENTAL_QUALITY, HISSE_RATIO)
    _apply_number_format_to_columns(ws, headers, value_columns, VALUE_NUMBER_FORMAT, first_data_row, last_data_row)
    _apply_number_format_to_columns(ws, headers, percent_columns, PERCENT_NUMBER_FORMAT, first_data_row, last_data_row)
    _apply_number_format_to_columns(ws, headers, (HISSE_MODEL_COUNT,), "0", first_data_row, last_data_row)

    affected_columns = _existing_column_indexes(
        headers,
        ("D12", HISSE_AVERAGE, HISSE_RATIO, HISSE_UPSIDE, HISSE_ANALYST_UPSIDE, HISSE_CONFIDENCE, HISSE_FUNDAMENTAL_QUALITY),
    )
    _remove_conditional_formatting_for_columns(ws, affected_columns)


def _apply_threshold_font_color(cell, value: Any, threshold: float) -> None:
    font = copy(cell.font)
    number = _safe_float(value)
    font.color = GREEN_FONT_COLOR if number is not None and number > threshold else RED_FONT_COLOR
    cell.font = font


def _apply_signal_font_colors(ws, columns: Sequence[str], header_row: int, row_count: int) -> None:
    headers = {column_name: index for index, column_name in enumerate(columns, start=1)}
    signal_columns = (
        (HISSE_RATIO, RASYO_PASS_THRESHOLD),
        (HISSE_UPSIDE, 0.0),
        (HISSE_ANALYST_UPSIDE, 0.0),
        (HISSE_CONFIDENCE, RASYO_PASS_THRESHOLD),
        (HISSE_FUNDAMENTAL_QUALITY, RASYO_PASS_THRESHOLD),
    )
    for row_index in range(header_row + 1, header_row + row_count + 1):
        for column_name, threshold in signal_columns:
            column_index = headers.get(column_name)
            if column_index is None:
                continue
            cell = ws.cell(row=row_index, column=column_index)
            _apply_threshold_font_color(cell, cell.value, threshold)


def _ensure_sheet(workbook, sheet_name: str):
    if sheet_name in workbook.sheetnames:
        return workbook[sheet_name]
    return workbook.create_sheet(sheet_name)


def _build_empty_report_workbook():
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name in OUTPUT_SHEETS:
        workbook.create_sheet(sheet_name)
    return workbook


def _load_or_create_report_workbook(template_path: Path):
    if template_path.exists():
        return load_workbook(template_path, data_only=False)
    return _build_empty_report_workbook()


def _keep_only_sheets(workbook, sheet_names: Sequence[str]) -> None:
    keep = set(sheet_names)
    for sheet_name in tuple(workbook.sheetnames):
        if sheet_name not in keep:
            del workbook[sheet_name]
    for target_index, sheet_name in enumerate(sheet_names):
        current_index = workbook.sheetnames.index(sheet_name)
        workbook.move_sheet(workbook[sheet_name], offset=target_index - current_index)


def _load_template_index_map(template_path: Path) -> dict[str, str]:
    if not template_path.exists():
        return {}
    workbook = load_workbook(template_path, data_only=False)
    ws = workbook[SHEET_ENDEKS]
    index_map: dict[str, str] = {}
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=3, values_only=True):
        code = row[0]
        index_name = row[2]
        if isinstance(code, str) and code.strip() and isinstance(index_name, str) and index_name.strip():
            index_map[code.strip().upper()] = index_name.strip()
    workbook.close()
    return index_map


def _current_xu100_set(template_path: Path) -> set[str]:
    base_map = _load_template_index_map(template_path)
    current = {code for code, index_name in base_map.items() if index_name == INDEX_XU100}
    current.update(SHEET_UPDATES_Q2_2026_XU100_ADD)
    current.difference_update(SHEET_UPDATES_Q2_2026_XU100_REMOVE)
    return current


def _sector_from_valuation_lookup(valuation_lookup: pd.DataFrame, code: str) -> str:
    if valuation_lookup.empty or code not in valuation_lookup.index:
        return ""
    return str(valuation_lookup.loc[code].get(SECTOR_COLUMN, "") or "")


def _ticker_index_name(ticker: str, code: str, xu100_codes: set[str]) -> str:
    if not ticker.upper().endswith(".IS"):
        return INDEX_US
    return INDEX_XU100 if code in xu100_codes else INDEX_XUTUM


def build_endeks_frame(
    tickers: list[str],
    valuation_frame: pd.DataFrame,
    template_path: Path,
) -> pd.DataFrame:
    xu100_codes = _current_xu100_set(template_path)
    valuation_lookup = valuation_frame.set_index(CODE_COLUMN) if not valuation_frame.empty and CODE_COLUMN in valuation_frame.columns else pd.DataFrame()

    rows: list[dict[str, Any]] = []
    for ticker in tickers:
        code = _normalized_code(ticker)
        sector_name = _sector_from_valuation_lookup(valuation_lookup, code)
        index_name = _ticker_index_name(ticker, code, xu100_codes)
        rows.append({CODE_COLUMN: code, SECTOR_COLUMN: sector_name, INDEX_COLUMN: index_name})
    return pd.DataFrame(rows, columns=_column_index((CODE_COLUMN, SECTOR_COLUMN, INDEX_COLUMN)))


def build_sector_frame(endeks_frame: pd.DataFrame, puan_frame: pd.DataFrame) -> pd.DataFrame:
    columns = (SECTOR_COLUMN, SECTOR_PB, SECTOR_PE, SECTOR_FD, SECTOR_EV_EBITDA)
    if endeks_frame.empty or puan_frame.empty:
        return pd.DataFrame(columns=_column_index(columns))

    merged = endeks_frame.merge(puan_frame, left_on=CODE_COLUMN, right_on=PUAN_CODE, how="left")
    rows: list[dict[str, Any]] = []
    for sector_name, sector_slice in merged.groupby(SECTOR_COLUMN, dropna=False):
        fd_series = (
            pd.to_numeric(sector_slice[PUAN_EV_EBITDA], errors="coerce")
            * pd.to_numeric(sector_slice[PUAN_EBITDA], errors="coerce")
        ).dropna()
        rows.append(
            {
                SECTOR_COLUMN: sector_name or "",
                SECTOR_PB: _median_or_none(sector_slice[PUAN_PB]),
                SECTOR_PE: _median_or_none(sector_slice[PUAN_PE]),
                SECTOR_FD: _median_or_none(fd_series),
                SECTOR_EV_EBITDA: _median_or_none(sector_slice[PUAN_EV_EBITDA]),
            }
        )
    return pd.DataFrame(rows, columns=_column_index(columns))


def _vars_values_for_market(config: dict[str, Any]) -> dict[str, float]:
    risk_free = float(config["risk_free_rate"])
    bond_yield_2 = float(config["bond_yield_2"])
    cost_of_debt = float(config["cost_of_debt"])
    sovereign_spread = max(bond_yield_2 - risk_free, 0.0)
    return {
        " 2Tahvil": bond_yield_2,
        " 5Tahvil": risk_free + sovereign_spread,
        "10Tahvil": risk_free,
        "K. Vergisi": float(config["tax_rate"]),
        "Faiz Oranı": cost_of_debt,
        "Risksiz Faiz": risk_free,
        "CDS 5(Ülke Riski)": sovereign_spread,
        "Borçlanma Maliyeti": cost_of_debt,
        "Piyasa Risk Primi": float(config["market_premium"]),
        "Terminal Büyüme": float(config["terminal_growth"]),
    }


def build_vars_frame() -> pd.DataFrame:
    config = load_macro_config()
    tr_values = _vars_values_for_market(config["tr"])
    us_values = _vars_values_for_market(config["us"])
    rows = [
        {"Tür": metric_name, TR_VALUE_COLUMN: tr_values[metric_name], US_VALUE_COLUMN: us_values[metric_name]}
        for metric_name in tr_values
    ]
    return pd.DataFrame(rows, columns=_column_index(("Tür", TR_VALUE_COLUMN, US_VALUE_COLUMN)))


def build_notes_frame(run_metadata: Optional[dict[str, Any]] = None) -> pd.DataFrame:
    rows: list[dict[str, Any]] = [
        {
            SECTION_COLUMN: "Metodoloji",
            TOPIC_COLUMN: HISSE_AVERAGE,
            VALUE_COLUMN: "D. Ort eşit ortalama değildir; sektör profiline, şirket kalitesine, çarpan riskine, veri tamlığına ve model sapmasına göre ayarlanmış ağırlıklı hedef değerdir.",
        },
        {
            SECTION_COLUMN: "Metodoloji",
            TOPIC_COLUMN: "Geçerlilik",
            VALUE_COLUMN: "Pozitif olmayan, boş veya güncel fiyatın %5'i ile 5 katı dışındaki model değerleri hedef fiyat hesabına ve rapora alınmaz; kalan ağırlıklar kendi içinde yeniden normalize edilir.",
        },
        {
            SECTION_COLUMN: "Metodoloji",
            TOPIC_COLUMN: HISSE_CONFIDENCE,
            VALUE_COLUMN: "Güven; geçerli model sayısı, modeller arası sapma, veri tamlığı, rasyo skoru, NIS işareti, gelir-varlık-borç-brüt marj kalitesi ve mevcut fiyat/analist hedefiyle tutarlılıktan oluşur. Aşırı FK/FD-FAVÖK ve negatif büyüme ilgili model ağırlığını düşürür.",
        },
        {
            SECTION_COLUMN: "Metodoloji",
            TOPIC_COLUMN: SHEET_KALITE,
            VALUE_COLUMN: "Kalite sayfası her D modeli için ham değer, filtre sonrası değer, ağırlık profili, kullanım durumu ve dışlanma gerekçesini gösterir. Bu sayfa denetim izi olarak kullanılmalıdır.",
        },
        {
            SECTION_COLUMN: "Metodoloji",
            TOPIC_COLUMN: WEIGHT_PROFILE_COLUMN,
            VALUE_COLUMN: "Ağırlık profili seçimleri repo kökündeki sector_profiles.json içindeki merkezi sektör→profil sözlüğü ve anahtar kelime fallback mantığıyla yapılır. Profil bazlı model ağırlıkları ve açıklamalar ise valuation_profiles.json dosyasından yüklenir.",
        },
        {
            SECTION_COLUMN: "Metodoloji",
            TOPIC_COLUMN: "Veri Sağlayıcı",
            VALUE_COLUMN: "Canlı veri zinciri taze cache → Yahoo Finance → Alpha Vantage (ALPHAVANTAGE_API_KEY tanımlıysa) → stale cache sırasıyla çalışır. Eksik kalan alanlar bir sonraki sağlayıcıdan güvenli backfill ile tamamlanır.",
        },
    ]
    for model_key in D_FIELDS:
        rows.append(
            {
                SECTION_COLUMN: "D Modeli",
                TOPIC_COLUMN: model_key,
                VALUE_COLUMN: HISSE_HEADER_COMMENTS.get(model_key, ""),
            }
        )
    for profile_name, weights in MODEL_WEIGHTS.items():
        for model_key in D_FIELDS:
            rows.append(
                {
                    SECTION_COLUMN: WEIGHT_COLUMN,
                    TOPIC_COLUMN: profile_name,
                    VALUE_COLUMN: model_key,
                    WEIGHT_COLUMN: weights.get(model_key, 0.0),
                }
            )
    for profile_name, note in PROFILE_NOTES.items():
        rows.append(
            {
                SECTION_COLUMN: "Profil",
                TOPIC_COLUMN: profile_name,
                VALUE_COLUMN: note,
            }
        )
    if isinstance(run_metadata, dict):
        for key in ("started_at", "input_path", "output_path", "tickers_total", "market_counts", "macro_config_as_of", "git_revision"):
            value = run_metadata.get(key)
            if value in (None, "", {}):
                continue
            rows.append(
                {
                    SECTION_COLUMN: "Çalıştırma",
                    TOPIC_COLUMN: str(key),
                    VALUE_COLUMN: str(value),
                }
            )
    return pd.DataFrame(rows, columns=_column_index((SECTION_COLUMN, TOPIC_COLUMN, VALUE_COLUMN, WEIGHT_COLUMN)))


def build_puan_frame(financials_frame: pd.DataFrame) -> pd.DataFrame:
    if financials_frame.empty:
        return pd.DataFrame(columns=_column_index(PUAN_COLUMNS))
    renamed = financials_frame.rename(
        columns={
            "Symbol": PUAN_CODE,
            "Name": PUAN_NAME,
            "EBITDA": PUAN_EBITDA,
            "ROE (%) Annual": PUAN_ROE,
            "EV/EBITDA": PUAN_EV_EBITDA,
            "P/CF": PUAN_PCF,
            "P/E": PUAN_PE,
            "EPS": PUAN_EPS,
            "Forward EPS": PUAN_FORWARD_EPS,
            "Forward P/E": PUAN_FORWARD_PE,
            "P/B": PUAN_PB,
            "P/S": PUAN_PS,
            "Beta": PUAN_BETA,
            "Asset Growth (%)": PUAN_ASSET_GROWTH,
            "Net Income Growth (Annual, %)": PUAN_NET_INCOME_GROWTH,
            "Net Income": PUAN_NET_INCOME,
            "Operating Income": PUAN_OPERATING_INCOME,
            "Shareholders' Equity": PUAN_EQUITY,
            "Debt-to-Assets Ratio": PUAN_DEBT_SOURCE,
            "Paid-in Capital": PUAN_PAID_IN_CAPITAL,
            "Net Working Capital": PUAN_NET_WORKING_CAPITAL,
            "Net Debt/EBITDA(Annual, %)": PUAN_NET_DEBT_TO_EBITDA,
            "Analyst Target Mean": PUAN_ANALYST_TARGET_MEAN,
            "Analyst Target Median": PUAN_ANALYST_TARGET_MEDIAN,
            "Analyst Count": PUAN_ANALYST_COUNT,
            "Earnings Growth": PUAN_EARNINGS_GROWTH,
            "Revenue Growth": PUAN_REVENUE_GROWTH,
            "Free Cash Flow": PUAN_FREE_CASH_FLOW,
            "Operating Cash Flow": PUAN_OPERATING_CASH_FLOW,
            "Total Revenue": PUAN_REVENUE,
            "Gross Margin (%)": PUAN_GROSS_MARGIN,
            "Total Assets": PUAN_TOTAL_ASSETS,
            "Total Debt": PUAN_TOTAL_DEBT,
        }
    )
    return renamed.reindex(columns=_column_index(PUAN_COLUMNS))


def _between(value: Any, lower: float, upper: float) -> Optional[bool]:
    number = _safe_float(value)
    if number is None:
        return None
    return lower < number < upper


def _greater_than(value: Any, threshold: float) -> Optional[bool]:
    number = _safe_float(value)
    if number is None:
        return None
    return number > threshold


def _clamp_score(value: float) -> float:
    return max(0.0, min(1.0, float(value)))


def _score_band(value: Optional[float], rule: dict[str, Any]) -> Optional[float]:
    if value is None:
        return None
    soft_min = _safe_float(rule.get("soft_min"))
    target_min = _safe_float(rule.get("target_min"))
    target_max = _safe_float(rule.get("target_max"))
    soft_max = _safe_float(rule.get("soft_max"))
    if None in (soft_min, target_min, target_max, soft_max):
        return None
    if value <= soft_min or value >= soft_max:
        return 0.0
    if target_min <= value <= target_max:
        return 1.0
    if value < target_min:
        return _clamp_score((value - soft_min) / (target_min - soft_min))
    return _clamp_score((soft_max - value) / (soft_max - target_max))


def _score_min(value: Optional[float], rule: dict[str, Any]) -> Optional[float]:
    if value is None:
        return None
    soft_min = _safe_float(rule.get("soft_min"))
    target_min = _safe_float(rule.get("target_min"))
    if None in (soft_min, target_min):
        return None
    if value <= soft_min:
        return 0.0
    if value >= target_min:
        return 1.0
    return _clamp_score((value - soft_min) / (target_min - soft_min))


def _metric_score(value: Optional[float], rule: dict[str, Any]) -> Optional[float]:
    kind = str(rule.get("kind", "") or "").strip()
    if kind == "band":
        return _score_band(value, rule)
    if kind == "min":
        return _score_min(value, rule)
    return None


def _ratio_metric_values(row: pd.Series) -> dict[str, Optional[float]]:
    return {
        "current_ratio": _safe_float(row.get(RASYO_CURRENT_RATIO)),
        "quick_ratio": _safe_float(row.get(RASYO_QUICK_RATIO)),
        "cash_ratio": _safe_float(row.get(RASYO_CASH_RATIO)),
        "roe": _safe_float(row.get(RASYO_ROE)),
        "pe": _safe_float(row.get(RASYO_PE)),
        "pb": _safe_float(row.get(RASYO_PB)),
        "ebitda_growth": _safe_float(row.get(RASYO_EBITDA_GROWTH)),
        "peg": _safe_float(row.get(RASYO_PEG)),
        "roic": _safe_float(row.get(RASYO_ROIC)),
        "asset_turnover": _safe_float(row.get(RASYO_ASSET_TURNOVER)),
    }


def _category_score_from_parts(parts: list[tuple[float, float]]) -> Optional[float]:
    total_weight = sum(weight for _, weight in parts)
    if total_weight <= 0:
        return None
    return _clamp_score(sum(score * weight for score, weight in parts) / total_weight)


def _empty_ratio_category_scores() -> dict[str, Optional[float]]:
    return dict.fromkeys(
        (
            RATIO_CATEGORY_LIQUIDITY,
            RATIO_CATEGORY_PROFITABILITY,
            RATIO_CATEGORY_VALUATION,
            RATIO_CATEGORY_GROWTH,
            RATIO_CATEGORY_EFFICIENCY,
        )
    )


def _base_ratio_category_parts() -> dict[str, list[tuple[float, float]]]:
    return {
        RATIO_CATEGORY_LIQUIDITY: [],
        RATIO_CATEGORY_PROFITABILITY: [],
        RATIO_CATEGORY_VALUATION: [],
        RATIO_CATEGORY_GROWTH: [],
        RATIO_CATEGORY_EFFICIENCY: [],
    }


def _ratio_profile_metrics(profile_name: str) -> tuple[dict[str, Any], Optional[float]]:
    profile = resolve_ratio_profile(RATIO_PROFILE_CONFIG, profile_name, "tr")
    if not isinstance(profile, dict):
        return {}, None
    metrics = profile.get("metrics", {})
    return (metrics if isinstance(metrics, dict) else {}), _safe_float(profile.get("missing_penalty"))


def _ratio_profile_metrics_for_market(profile_name: str, market_key: str) -> tuple[dict[str, Any], Optional[float]]:
    profile = resolve_ratio_profile(RATIO_PROFILE_CONFIG, profile_name, market_key)
    if not isinstance(profile, dict):
        return {}, None
    metrics = profile.get("metrics", {})
    return (metrics if isinstance(metrics, dict) else {}), _safe_float(profile.get("missing_penalty"))


def _accumulate_ratio_metric_scores(
    metrics: dict[str, Any],
    metric_values: dict[str, Optional[float]],
    category_parts: dict[str, list[tuple[float, float]]],
) -> tuple[float, float, float]:
    total_weight = 0.0
    present_weight = 0.0
    weighted_score = 0.0
    for metric_name, rule in metrics.items():
        if not isinstance(rule, dict):
            continue
        weight = _safe_float(rule.get("weight")) or 0.0
        if weight <= 0:
            continue
        total_weight += weight
        score = _metric_score(metric_values.get(metric_name), rule)
        if score is None:
            continue
        present_weight += weight
        weighted_score += score * weight
        category_name = str(rule.get("category", "") or "").strip()
        if category_name in category_parts:
            category_parts[category_name].append((score, weight))
    return total_weight, present_weight, weighted_score


def _ratio_score_summary(row: pd.Series, profile_name: str, market_key: str) -> RatioScoreSummary:
    metrics, missing_penalty_factor = _ratio_profile_metrics_for_market(profile_name, market_key)
    metric_values = _ratio_metric_values(row)
    category_parts = _base_ratio_category_parts()
    total_weight, present_weight, weighted_score = _accumulate_ratio_metric_scores(
        metrics,
        metric_values,
        category_parts,
    )

    if total_weight <= 0 or present_weight <= 0:
        return RatioScoreSummary(
            profile_name=profile_name,
            normalized_score=None,
            score_100=None,
            coverage=None,
            missing_penalty=None,
            category_scores=_empty_ratio_category_scores(),
        )

    coverage = _clamp_score(present_weight / total_weight)
    present_average = weighted_score / present_weight
    missing_penalty = (1.0 - coverage) * (missing_penalty_factor if missing_penalty_factor is not None else 0.15)
    normalized_score = _clamp_score((present_average * coverage) + (0.50 * (1.0 - coverage)) - missing_penalty)
    category_scores = {
        category: _category_score_from_parts(parts)
        for category, parts in category_parts.items()
    }
    return RatioScoreSummary(
        profile_name=profile_name,
        normalized_score=normalized_score,
        score_100=round(normalized_score * 100.0, 1),
        coverage=coverage,
        missing_penalty=missing_penalty,
        category_scores=category_scores,
    )


def _sector_lookup_for_ratio(endeks_frame: pd.DataFrame) -> dict[str, str]:
    if endeks_frame.empty or CODE_COLUMN not in endeks_frame.columns or SECTOR_COLUMN not in endeks_frame.columns:
        return {}
    lookup: dict[str, str] = {}
    for _, row in endeks_frame.iterrows():
        code = _normalized_code(str(row.get(CODE_COLUMN, "") or ""))
        sector_name = str(row.get(SECTOR_COLUMN, "") or "")
        if code:
            lookup[code] = sector_name
    return lookup


def _market_key_from_symbol(symbol: str) -> str:
    return "tr" if str(symbol or "").upper().endswith(".IS") else "us"


def _quality_status(
    filtered_value: Optional[float],
    adjusted_weight: float,
    model_key: str,
    retained_keys: set[str],
) -> str:
    del model_key, retained_keys
    if filtered_value is None:
        return QUALITY_STATUS_FILTERED
    if adjusted_weight <= 0:
        return QUALITY_STATUS_UNWEIGHTED
    return QUALITY_STATUS_USED


def _quality_row(
    inputs: HisseInputs,
    model_key: str,
    raw_values: dict[str, Optional[float]],
    filtered_values: dict[str, Optional[float]],
    weight_profile: str,
    base_weights: dict[str, float],
    adjusted_weights: dict[str, float],
    retained_keys: set[str],
    quality_context: QualityRowContext,
) -> dict[str, Any]:
    filtered_value = filtered_values.get(model_key)
    adjusted_weight = adjusted_weights.get(model_key, 0.0)
    included_in_fair_value = model_key in retained_keys and adjusted_weight > 0
    status = _quality_status(filtered_value, adjusted_weight, model_key, retained_keys)
    fundamental_score = _fundamental_quality_score(inputs)
    return {
        QUALITY_CODE: inputs.code,
        QUALITY_MODEL: model_key,
        QUALITY_RAW_VALUE: _round_or_none(raw_values.get(model_key)),
        QUALITY_FILTERED_VALUE: _round_or_none(filtered_value),
        QUALITY_STATUS: status,
        QUALITY_REASON: _model_value_reason(raw_values.get(model_key), filtered_value, inputs.price),
        QUALITY_WEIGHT_PROFILE: weight_profile,
        QUALITY_BASE_WEIGHT: _round_or_none(base_weights.get(model_key, 0.0), 4),
        QUALITY_ADJUSTED_WEIGHT: _round_or_none(adjusted_weight, 4),
        QUALITY_INCLUDED_IN_FAIR_VALUE: "Evet" if included_in_fair_value else "Hayır",
        QUALITY_PRICE: _round_or_none(inputs.price),
        QUALITY_SECTOR: inputs.sector_name,
        QUALITY_INDEX: inputs.index_name,
        QUALITY_SIGNAL_SCORE: _round_or_none(quality_context.signal_score, 0),
        QUALITY_SIGNAL_LABEL: quality_context.signal_label,
        QUALITY_WARNING_COUNT: _round_or_none(quality_context.warning_count, 0),
        QUALITY_MODEL_WARNINGS: quality_context.model_warnings,
        QUALITY_DATA_COMPLETENESS: _round_or_none(quality_context.data_completeness, 4),
        QUALITY_FUNDAMENTAL_SCORE: _round_or_none(fundamental_score, 4),
        QUALITY_MODEL_COUNT: len(adjusted_weights),
        QUALITY_RESULT_STATUS: quality_context.summary.status,
        QUALITY_RESULT_NOTE: quality_context.summary.note,
    }


def build_rasyo_frame(ratio_frame: pd.DataFrame, endeks_frame: pd.DataFrame) -> pd.DataFrame:
    if ratio_frame.empty:
        return pd.DataFrame(columns=_column_index(RASYO_COLUMNS))
    renamed = ratio_frame.rename(
        columns={
            "F/K": RASYO_PE,
            HISSE_PE: RASYO_PE,
            "Özsermaye Kârlılığı (ROE) (%) Yıllık": RASYO_ROE,
            HISSE_PB: RASYO_PB,
            "PEG Oranı": RASYO_PEG,
            "ROIC (%)": RASYO_ROIC,
            RASYO_STOCK: RASYO_STOCK,
        }
    )
    sector_lookup = _sector_lookup_for_ratio(endeks_frame)
    summaries: list[RatioScoreSummary] = []
    for _, row in renamed.iterrows():
        raw_symbol = str(row.get(RASYO_STOCK, "") or "")
        code = _normalized_code(raw_symbol)
        sector_name = sector_lookup.get(code, "")
        market_key = _market_key_from_symbol(raw_symbol)
        if not sector_name and bool(row.get("Finansal Sektor")):
            sector_name = "Banka"
        profile_name = _weight_profile_for_sector(sector_name)
        summaries.append(_ratio_score_summary(row, profile_name, market_key))

    renamed[RASYO_PROFILE] = [summary.profile_name for summary in summaries]
    renamed[RASYO_SCORE_100] = [_round_or_none(summary.score_100, 1) for summary in summaries]
    renamed[RASYO_SCORE] = [_round_or_none(summary.normalized_score, 4) for summary in summaries]
    renamed[RASYO_COVERAGE] = [_round_or_none(summary.coverage, 4) for summary in summaries]
    renamed[RASYO_MISSING_PENALTY] = [_round_or_none(summary.missing_penalty, 4) for summary in summaries]
    renamed[RASYO_LIQUIDITY_SCORE] = [_round_or_none(summary.category_scores.get("Likidite"), 4) for summary in summaries]
    renamed[RASYO_PROFITABILITY_SCORE] = [_round_or_none(summary.category_scores.get("Karlılık"), 4) for summary in summaries]
    renamed[RASYO_VALUATION_SCORE] = [_round_or_none(summary.category_scores.get("Değerleme"), 4) for summary in summaries]
    renamed[RASYO_GROWTH_SCORE] = [_round_or_none(summary.category_scores.get("Büyüme"), 4) for summary in summaries]
    renamed[RASYO_EFFICIENCY_SCORE] = [_round_or_none(summary.category_scores.get("Verimlilik"), 4) for summary in summaries]
    return renamed.reindex(columns=_column_index(RASYO_COLUMNS))


def _macro_rates_from_config(market_key: str, macro_config: dict[str, Any]) -> MacroRates:
    return MacroRates(
        market_key=market_key,
        two_year_bond=_rate_decimal(macro_config.get("bond_yield_2")),
        market_premium=_rate_decimal(macro_config.get("market_premium")),
        terminal_growth=_rate_decimal(macro_config.get("terminal_growth")),
        cost_of_debt=_rate_decimal(macro_config.get("cost_of_debt")),
        tax_rate=_rate_decimal(macro_config.get("tax_rate")),
        risk_free_rate=_rate_decimal(macro_config.get("risk_free_rate")),
    )


def _load_macro_rate_book() -> MacroRateBook:
    config = load_macro_config()
    return MacroRateBook(
        tr=_macro_rates_from_config("tr", config["tr"]),
        us=_macro_rates_from_config("us", config["us"]),
    )


def _empty_series() -> pd.Series:
    return pd.Series(dtype="object")


def _lookup_row(frame: pd.DataFrame, key: str) -> pd.Series:
    if frame.empty or key not in frame.index:
        return _empty_series()
    row = frame.loc[key]
    if isinstance(row, pd.DataFrame):
        return row.iloc[-1]
    return row


def _build_rasyo_lookup(rasyo_frame: pd.DataFrame) -> pd.DataFrame:
    if rasyo_frame.empty or RASYO_STOCK not in rasyo_frame.columns:
        return pd.DataFrame()
    rasyo_indexed = rasyo_frame.copy()
    rasyo_indexed["_Kod"] = rasyo_indexed[RASYO_STOCK].apply(_normalized_code)
    return rasyo_indexed.drop_duplicates(subset=["_Kod"], keep="last").set_index("_Kod")


def _build_row_lookups(
    valuation_frame: pd.DataFrame,
    puan_frame: pd.DataFrame,
    rasyo_frame: pd.DataFrame,
    endeks_frame: pd.DataFrame,
    sektor_frame: pd.DataFrame,
    ina_frame: pd.DataFrame,
) -> RowLookups:
    return RowLookups(
        valuation=valuation_frame.set_index(CODE_COLUMN) if not valuation_frame.empty else pd.DataFrame(),
        puan=puan_frame.set_index(PUAN_CODE) if not puan_frame.empty else pd.DataFrame(),
        rasyo=_build_rasyo_lookup(rasyo_frame),
        endeks=endeks_frame.set_index(CODE_COLUMN) if not endeks_frame.empty else pd.DataFrame(),
        sektor=sektor_frame.set_index(SECTOR_COLUMN) if not sektor_frame.empty else pd.DataFrame(),
        ina=ina_frame.set_index(CODE_COLUMN) if not ina_frame.empty and CODE_COLUMN in ina_frame.columns else pd.DataFrame(),
    )


def _nis_sign(value: Any) -> str:
    nis_value = _safe_float(value)
    if nis_value is None or nis_value == 0:
        return ""
    return "+" if nis_value > 0 else "-"


def _net_debt_from_ratio(ebitda: Optional[float], net_debt_to_ebitda: Optional[float]) -> float:
    if ebitda is None or net_debt_to_ebitda is None:
        return 0.0
    return ebitda * net_debt_to_ebitda


def _hisse_inputs(ticker: str, lookups: RowLookups) -> HisseInputs:
    code = _normalized_code(ticker)
    valuation_row = _lookup_row(lookups.valuation, code)
    puan_row = _lookup_row(lookups.puan, code)
    rasyo_row = _lookup_row(lookups.rasyo, code)
    endeks_row = _lookup_row(lookups.endeks, code)
    ina_row = _lookup_row(lookups.ina, code)
    sector_name = str(endeks_row.get(SECTOR_COLUMN, valuation_row.get(SECTOR_COLUMN, "")) or "")
    sektor_row = _lookup_row(lookups.sektor, sector_name)
    ebitda = _safe_float(puan_row.get(PUAN_EBITDA))
    net_debt_to_ebitda = _rate_decimal(puan_row.get(PUAN_NET_DEBT_TO_EBITDA))
    index_default = INDEX_US if not ticker.upper().endswith(".IS") else INDEX_XUTUM
    return HisseInputs(
        code=code,
        sector_name=sector_name,
        index_name=str(endeks_row.get(INDEX_COLUMN, index_default) or index_default),
        market_key="tr" if ticker.upper().endswith(".IS") else "us",
        price=_safe_float(valuation_row.get("Güncel Fiyat")),
        company_pe=_safe_float(puan_row.get(PUAN_PE)),
        company_pb=_safe_float(puan_row.get(PUAN_PB)),
        nis_sign=_nis_sign(puan_row.get(PUAN_NET_WORKING_CAPITAL)),
        ratio_score=_safe_float(rasyo_row.get(RASYO_SCORE)),
        sector_pe=_safe_float(sektor_row.get(SECTOR_PE)),
        sector_pb=_safe_float(sektor_row.get(SECTOR_PB)),
        sector_ev_ebitda=_safe_float(sektor_row.get(SECTOR_EV_EBITDA)),
        eps=_safe_float(puan_row.get(PUAN_EPS)),
        forward_eps=_safe_float(puan_row.get(PUAN_FORWARD_EPS)),
        forward_pe=_safe_float(puan_row.get(PUAN_FORWARD_PE)),
        roe=_rate_decimal(puan_row.get(PUAN_ROE)),
        beta=_safe_float(puan_row.get(PUAN_BETA)),
        debt_ratio=_debt_weight_from_debt_source(puan_row.get(PUAN_DEBT_SOURCE)),
        net_income_growth=_rate_decimal(puan_row.get(PUAN_NET_INCOME_GROWTH)),
        earnings_growth=_rate_decimal(puan_row.get(PUAN_EARNINGS_GROWTH)),
        revenue_growth=_rate_decimal(puan_row.get(PUAN_REVENUE_GROWTH)),
        net_income=_safe_float(puan_row.get(PUAN_NET_INCOME)),
        operating_income=_safe_float(puan_row.get(PUAN_OPERATING_INCOME)),
        equity=_safe_float(puan_row.get(PUAN_EQUITY)),
        paid_in_capital=_safe_float(puan_row.get(PUAN_PAID_IN_CAPITAL)),
        ebitda=ebitda,
        free_cash_flow=_safe_float(puan_row.get(PUAN_FREE_CASH_FLOW)),
        operating_cash_flow=_safe_float(puan_row.get(PUAN_OPERATING_CASH_FLOW)),
        net_debt=_net_debt_from_ratio(ebitda, net_debt_to_ebitda),
        asset_growth=_safe_float(puan_row.get(PUAN_ASSET_GROWTH)),
        dcf_value=_positive_or_none(ina_row.get(INA_VALUE_COLUMN)),
        analyst_target=_safe_float(puan_row.get(PUAN_ANALYST_TARGET_MEAN)) or _safe_float(puan_row.get(PUAN_ANALYST_TARGET_MEDIAN)),
        analyst_count=_safe_float(puan_row.get(PUAN_ANALYST_COUNT)),
        revenue=_safe_float(puan_row.get(PUAN_REVENUE)),
        gross_margin=_safe_float(puan_row.get(PUAN_GROSS_MARGIN)),
        total_assets=_safe_float(puan_row.get(PUAN_TOTAL_ASSETS)),
        total_debt=_safe_float(puan_row.get(PUAN_TOTAL_DEBT)),
    )


def _terminal_value_per_share(
    base_value: Optional[float],
    paid_in_capital: Optional[float],
    discount_rate: Optional[float],
    terminal_growth: Optional[float],
) -> Optional[float]:
    if discount_rate is None or terminal_growth is None:
        return None
    numerator = base_value * (1 + terminal_growth) if base_value is not None else None
    total_value = _safe_divide(numerator, discount_rate - terminal_growth)
    return _safe_divide(total_value, paid_in_capital)


def _raw_valuation_points(inputs: HisseInputs, macro_rates: MacroRates) -> dict[str, Optional[float]]:
    book_value_per_share = _safe_divide(inputs.equity, inputs.paid_in_capital)
    enterprise_value = None
    if inputs.ebitda is not None and inputs.sector_ev_ebitda is not None:
        enterprise_value = inputs.ebitda * inputs.sector_ev_ebitda
    discount_rate = _discount_rate(
        debt_ratio=inputs.debt_ratio,
        beta=inputs.beta,
        cost_of_debt=macro_rates.cost_of_debt,
        tax_rate=macro_rates.tax_rate,
        risk_free_rate=macro_rates.risk_free_rate,
        market_premium=macro_rates.market_premium,
    )
    cost_of_equity = _cost_of_equity(
        beta=inputs.beta,
        risk_free_rate=macro_rates.risk_free_rate,
        market_premium=macro_rates.market_premium,
    )
    d6 = None
    if inputs.forward_eps is not None and inputs.forward_pe is not None:
        d6 = inputs.forward_eps * inputs.forward_pe
    elif inputs.eps is not None and inputs.asset_growth is not None and inputs.company_pe is not None:
        d6 = inputs.eps * (1 + inputs.asset_growth / 100.0) * inputs.company_pe
    d1 = None
    if inputs.sector_pe is not None and inputs.eps is not None:
        d1 = inputs.sector_pe * inputs.eps
    d2_numerator = inputs.operating_income * 10 if inputs.operating_income is not None else None
    d3_numerator = inputs.net_income * 10 if inputs.net_income is not None else None
    d4 = None
    if book_value_per_share is not None and inputs.sector_pb is not None:
        d4 = book_value_per_share * inputs.sector_pb
    d7_numerator = enterprise_value - inputs.net_debt if enterprise_value is not None else None
    d8_numerator = None
    if inputs.eps is not None and inputs.sector_pe is not None:
        d8_numerator = inputs.eps * inputs.sector_pe
    return {
        "D1": d1,
        "D2": _safe_divide(d2_numerator, inputs.paid_in_capital),
        "D3": _safe_divide(d3_numerator, inputs.paid_in_capital),
        "D4": d4,
        "D5": book_value_per_share,
        "D6": d6,
        "D7": _safe_divide(d7_numerator, inputs.paid_in_capital),
        "D8": _bond_adjusted_multiple_value(d8_numerator, macro_rates.two_year_bond, inputs.market_key),
        "D9": _terminal_value_per_share(inputs.operating_income, inputs.paid_in_capital, discount_rate, macro_rates.terminal_growth),
        "D10": _terminal_value_per_share(inputs.net_income, inputs.paid_in_capital, discount_rate, macro_rates.terminal_growth),
        "D11": inputs.dcf_value,
        "D12": _roe_justified_pb_value(
            book_value_per_share=book_value_per_share,
            roe=inputs.roe,
            cost_of_equity=cost_of_equity,
            terminal_growth=macro_rates.terminal_growth,
        ),
    }


def _valuation_points(inputs: HisseInputs, macro_rates: MacroRates) -> dict[str, Optional[float]]:
    raw_values = _raw_valuation_points(inputs, macro_rates)
    return sanity_checked_valuation_points(raw_values, inputs.price)


def _hisse_output_row(inputs: HisseInputs, values: dict[str, Optional[float]]) -> dict[str, Any]:
    weight_profile = _weight_profile_for_sector(inputs.sector_name)
    summary = valuation_summary(values, inputs, weight_profile)
    fundamental_score = _fundamental_quality_score(inputs)
    avg_fair = summary.fair_value
    gp_pct = None
    if avg_fair is not None and inputs.price not in (None, 0):
        gp_pct = (avg_fair - inputs.price) / inputs.price
    analyst_gp_pct = None
    if inputs.analyst_target is not None and inputs.price not in (None, 0):
        analyst_gp_pct = (inputs.analyst_target - inputs.price) / inputs.price
    return {
        CODE_COLUMN: inputs.code,
        SECTOR_COLUMN: inputs.sector_name,
        INDEX_COLUMN: inputs.index_name,
        HISSE_PRICE: inputs.price,
        HISSE_PE: inputs.company_pe,
        HISSE_PB: inputs.company_pb,
        HISSE_NIS: inputs.nis_sign,
        HISSE_RATIO: inputs.ratio_score,
        **values,
        HISSE_AVERAGE: _round_or_none(avg_fair),
        HISSE_ANALYST_TARGET: inputs.analyst_target,
        HISSE_UPSIDE: _round_or_none(gp_pct, 4),
        HISSE_ANALYST_UPSIDE: _round_or_none(analyst_gp_pct, 4),
        HISSE_CONFIDENCE: _round_or_none(summary.confidence, 4),
        HISSE_FUNDAMENTAL_QUALITY: _round_or_none(fundamental_score, 4),
        HISSE_MODEL_COUNT: summary.model_count,
        HISSE_STATUS: summary.status,
        HISSE_STATUS_NOTE: summary.note,
    }


def build_hisse_frame(
    tickers: list[str],
    valuation_frame: pd.DataFrame,
    puan_frame: pd.DataFrame,
    rasyo_frame: pd.DataFrame,
    endeks_frame: pd.DataFrame,
    sektor_frame: pd.DataFrame,
    ina_frame: pd.DataFrame,
) -> pd.DataFrame:
    lookups = _build_row_lookups(valuation_frame, puan_frame, rasyo_frame, endeks_frame, sektor_frame, ina_frame)
    macro_rate_book = _load_macro_rate_book()
    rows = []
    for ticker in tickers:
        inputs = _hisse_inputs(ticker, lookups)
        macro_rates = macro_rate_book.for_ticker(ticker)
        values = _valuation_points(inputs, macro_rates)
        rows.append(_hisse_output_row(inputs, values))
    return pd.DataFrame(rows, columns=_column_index(HISSE_COLUMNS))


def build_quality_frame(
    tickers: list[str],
    valuation_frame: pd.DataFrame,
    puan_frame: pd.DataFrame,
    rasyo_frame: pd.DataFrame,
    endeks_frame: pd.DataFrame,
    sektor_frame: pd.DataFrame,
    ina_frame: pd.DataFrame,
) -> pd.DataFrame:
    lookups = _build_row_lookups(valuation_frame, puan_frame, rasyo_frame, endeks_frame, sektor_frame, ina_frame)
    macro_rate_book = _load_macro_rate_book()
    valuation_lookup = valuation_frame.set_index(CODE_COLUMN) if not valuation_frame.empty and CODE_COLUMN in valuation_frame.columns else pd.DataFrame()
    rows: list[dict[str, Any]] = []

    for ticker in tickers:
        inputs = _hisse_inputs(ticker, lookups)
        macro_rates = macro_rate_book.for_ticker(ticker)
        raw_values = _raw_valuation_points(inputs, macro_rates)
        filtered_values = sanity_checked_valuation_points(raw_values, inputs.price)
        weight_profile = _weight_profile_for_sector(inputs.sector_name)
        summary = valuation_summary(filtered_values, inputs, weight_profile)
        base_weights = resolve_valuation_weight_map(VALUATION_PROFILE_CONFIG, weight_profile, inputs.market_key)
        adjusted_weights = _adjusted_model_weights(filtered_values, inputs, weight_profile)
        retained_keys = set(adjusted_weights)
        valuation_row = _lookup_row(valuation_lookup, inputs.code)
        signal_score = _safe_float(valuation_row.get(QUALITY_SIGNAL_SCORE))
        signal_label = str(valuation_row.get(QUALITY_SIGNAL_LABEL, "") or "")
        warning_count = _safe_float(valuation_row.get(QUALITY_WARNING_COUNT))
        model_warnings = str(valuation_row.get(QUALITY_MODEL_WARNINGS, "") or "")
        data_completeness = _data_completeness_score(inputs)
        quality_context = QualityRowContext(
            signal_score=signal_score,
            signal_label=signal_label,
            warning_count=warning_count,
            model_warnings=model_warnings,
            data_completeness=data_completeness,
            summary=summary,
        )

        for model_key in D_FIELDS:
            rows.append(
                _quality_row(
                    inputs,
                    model_key,
                    raw_values,
                    filtered_values,
                    weight_profile,
                    base_weights,
                    adjusted_weights,
                    retained_keys,
                    quality_context,
                )
            )
    return pd.DataFrame(rows, columns=_column_index(QUALITY_COLUMNS))


def write_template_report(
    output_path: str,
    tickers: list[str],
    valuation_frame: pd.DataFrame,
    financials_frame: pd.DataFrame,
    dcf_frame: pd.DataFrame,
    ratio_frame: pd.DataFrame,
    template_path: Path = DEFAULT_TEMPLATE_PATH,
    run_metadata: Optional[dict[str, Any]] = None,
) -> None:
    workbook = _load_or_create_report_workbook(template_path)

    endeks_frame = build_endeks_frame(tickers, valuation_frame, template_path)
    puan_frame = build_puan_frame(financials_frame)
    sektor_frame = build_sector_frame(endeks_frame, puan_frame)
    vars_frame = build_vars_frame()
    rasyo_frame = build_rasyo_frame(ratio_frame, endeks_frame)
    notes_frame = build_notes_frame(run_metadata)
    ina_frame = dcf_frame.copy()
    if not ina_frame.empty and INA_VALUE_COLUMN not in ina_frame.columns:
        ina_frame[INA_VALUE_COLUMN] = pd.NA
    hisse_frame = build_hisse_frame(tickers, valuation_frame, puan_frame, rasyo_frame, endeks_frame, sektor_frame, ina_frame)
    quality_frame = build_quality_frame(tickers, valuation_frame, puan_frame, rasyo_frame, endeks_frame, sektor_frame, ina_frame)

    ws_endeks = _ensure_sheet(workbook, SHEET_ENDEKS)
    _unmerge_sheet(ws_endeks)
    _clear_sheet(ws_endeks, start_row=1, start_col=1, end_col=3)
    _write_table(ws_endeks, endeks_frame, header_row=1, include_header=False)

    ws_sektor = _ensure_sheet(workbook, SHEET_SEKTOR)
    _unmerge_sheet(ws_sektor)
    _clear_sheet(ws_sektor, start_row=1, start_col=1, end_col=6)
    _write_table(ws_sektor, sektor_frame, header_row=1)

    ws_vars = _ensure_sheet(workbook, SHEET_VARS)
    _unmerge_sheet(ws_vars)
    _clear_sheet(ws_vars, start_row=1, start_col=1, end_col=3)
    _write_table(ws_vars, vars_frame, header_row=1)

    ws_rasyo = _ensure_sheet(workbook, SHEET_RASYO)
    _unmerge_sheet(ws_rasyo)
    _clear_sheet(ws_rasyo, start_row=1, start_col=1, end_col=len(RASYO_COLUMNS))
    _write_table(ws_rasyo, rasyo_frame, header_row=1)

    ws_puan = _ensure_sheet(workbook, SHEET_PUAN)
    _unmerge_sheet(ws_puan)
    _clear_sheet(ws_puan, start_row=1, start_col=1, end_col=len(PUAN_COLUMNS))
    _write_table(ws_puan, puan_frame, header_row=1)

    ws_notlar = _ensure_sheet(workbook, SHEET_NOTLAR)
    _unmerge_sheet(ws_notlar)
    _clear_sheet(ws_notlar, start_row=1, start_col=1)
    _write_table(ws_notlar, notes_frame, header_row=1)
    for row_index in range(2, len(notes_frame) + 2):
        ws_notlar.cell(row=row_index, column=4).number_format = PERCENT_NUMBER_FORMAT

    ws_kalite = _ensure_sheet(workbook, SHEET_KALITE)
    _unmerge_sheet(ws_kalite)
    _clear_sheet(ws_kalite, start_row=1, start_col=1, end_col=len(QUALITY_COLUMNS))
    _write_table(ws_kalite, quality_frame, header_row=1)

    ws_hisse = _ensure_sheet(workbook, SHEET_HISSE)
    _unmerge_sheet(ws_hisse)
    _clear_sheet(ws_hisse, start_row=2, start_col=1, end_col=len(HISSE_COLUMNS))
    _write_table(ws_hisse, hisse_frame, header_row=2)
    _sync_excel_table(ws_hisse, "Hisseler", HISSE_COLUMNS, header_row=2, row_count=len(hisse_frame))
    _align_hisse_column_styles(ws_hisse, HISSE_COLUMNS, header_row=2, row_count=len(hisse_frame))
    _apply_header_comments(ws_hisse, HISSE_COLUMNS, HISSE_HEADER_COMMENTS, header_row=2)
    _apply_hisse_number_formats(ws_hisse, HISSE_COLUMNS, header_row=2, row_count=len(hisse_frame))
    _apply_signal_font_colors(ws_hisse, HISSE_COLUMNS, header_row=2, row_count=len(hisse_frame))

    _keep_only_sheets(workbook, OUTPUT_SHEETS)

    workbook.save(output_path)
    workbook.close()
