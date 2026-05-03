import math

import pandas as pd
from openpyxl import load_workbook

from data.market_data_provider import TickerBundle
from reporting.financials_report import _financials_row_from_bundle, compute_share_count
from data.ratio_profile_config import load_ratio_profile_config, resolve_ratio_profile
from data.sector import US_PROFILE
from data.sector_profile_config import load_sector_profile_config
from data.valuation_profile_config import load_valuation_profile_config, resolve_valuation_weight_map
from reporting.template_report import (
    CODE_COLUMN,
    HisseInputs,
    INA_VALUE_COLUMN,
    MacroRates,
    RASYO_COVERAGE,
    RASYO_PROFILE,
    RASYO_SCORE,
    RASYO_SCORE_100,
    RASYO_VALUATION_SCORE,
    SHEET_KALITE,
    PUAN_GROSS_MARGIN,
    PUAN_REVENUE,
    PUAN_TOTAL_ASSETS,
    PUAN_TOTAL_DEBT,
    HISSE_FUNDAMENTAL_QUALITY,
    WEIGHT_PROFILE_DEFAULT,
    WEIGHT_PROFILE_FINANCIAL,
    WEIGHT_PROFILE_ENERGY_EQUIPMENT,
    WEIGHT_PROFILE_GROWTH,
    build_puan_frame,
    build_rasyo_frame,
    _hisse_output_row,
    _weight_profile_for_sector,
    _model_value_reason,
    _raw_valuation_points,
    sanity_checked_valuation_points,
    valuation_summary,
    write_template_report,
)


def test_financials_row_includes_assets_revenue_gross_margin_and_total_debt() -> None:
    bundle = TickerBundle(
        ticker="TEST.IS",
        info={
            "longName": "Test Sirket",
            "sharesOutstanding": 10.0,
            "trailingPE": 12.0,
            "priceToBook": 1.5,
            "beta": 1.0,
        },
        fast_info={},
        financials=pd.DataFrame(
            {
                pd.Timestamp("2025-12-31"): {
                    "Total Revenue": 1000.0,
                    "Gross Profit": 400.0,
                    "Net Income": 120.0,
                    "Operating Income": 160.0,
                }
            }
        ),
        balance_sheet=pd.DataFrame(
            {
                pd.Timestamp("2025-12-31"): {
                    "Total Assets": 2000.0,
                    "Total Debt": 600.0,
                    "Stockholders Equity": 900.0,
                    "Current Assets": 700.0,
                    "Current Liabilities": 300.0,
                    "Cash And Cash Equivalents": 100.0,
                    "Ordinary Shares Number": 10.0,
                }
            }
        ),
        cashflow=pd.DataFrame(),
        last_close=None,
        warnings=[],
    )

    row = _financials_row_from_bundle("TEST.IS", bundle)

    assert math.isclose(row["Total Revenue"], 1000.0)
    assert math.isclose(row["Total Assets"], 2000.0)
    assert math.isclose(row["Total Debt"], 600.0)
    assert math.isclose(row["Gross Margin (%)"], 40.0)


def test_build_puan_frame_maps_new_financial_signal_columns() -> None:
    frame = pd.DataFrame(
        [
            {
                "Symbol": "TEST",
                "Total Revenue": 1000.0,
                "Gross Margin (%)": 40.0,
                "Total Assets": 2000.0,
                "Total Debt": 600.0,
            }
        ]
    )

    result = build_puan_frame(frame)
    row = result.iloc[0]

    assert math.isclose(row[PUAN_REVENUE], 1000.0)
    assert math.isclose(row[PUAN_TOTAL_ASSETS], 2000.0)
    assert math.isclose(row[PUAN_TOTAL_DEBT], 600.0)
    assert math.isclose(row[PUAN_GROSS_MARGIN], 40.0)


def test_share_count_prefers_implied_shares_when_balance_stock_is_nominal() -> None:
    info = {
        "sharesOutstanding": None,
        "floatShares": None,
        "impliedShares": 3_750_000_000,
    }

    result = compute_share_count(info, annual_balance=pd.DataFrame())

    assert result == 3_750_000_000


def test_valuation_points_drop_extreme_values_against_price() -> None:
    values = {
        "D1": 100.0,
        "D2": 3.0,
        "D3": 700.0,
        "D4": -20.0,
    }

    result = sanity_checked_valuation_points(values, price=100.0)

    assert result == {
        "D1": 100.0,
        "D2": None,
        "D3": None,
        "D4": None,
    }


def test_high_forward_pe_reduces_growth_model_weight() -> None:
    inputs = HisseInputs(
        code="TEST",
        sector_name="Teknoloji",
        index_name="US",
        market_key="us",
        price=100.0,
        company_pe=150.0,
        company_pb=10.0,
        nis_sign="+",
        ratio_score=0.80,
        sector_pe=150.0,
        sector_pb=10.0,
        sector_ev_ebitda=20.0,
        eps=1.0,
        forward_eps=2.0,
        forward_pe=150.0,
        roe=0.20,
        beta=1.2,
        debt_ratio=0.10,
        net_income_growth=0.10,
        earnings_growth=0.20,
        revenue_growth=0.10,
        net_income=100.0,
        operating_income=120.0,
        equity=500.0,
        paid_in_capital=10.0,
        ebitda=150.0,
        free_cash_flow=90.0,
        operating_cash_flow=110.0,
        net_debt=20.0,
        asset_growth=10.0,
        dcf_value=100.0,
        analyst_target=110.0,
        analyst_count=20.0,
    )

    summary = valuation_summary(
        {"D1": 100.0, "D6": 400.0, "D7": 100.0, "D11": 100.0},
        inputs,
        WEIGHT_PROFILE_GROWTH,
    )

    assert summary.fair_value is not None
    assert summary.fair_value < 180.0
    assert summary.confidence is not None
    assert 0 < summary.confidence <= 1
    assert summary.publishable is True


def test_write_template_report_creates_workbook_when_template_is_missing(tmp_path) -> None:
    output_path = tmp_path / "report.xlsx"
    missing_template = tmp_path / "missing_template.xlsx"

    valuation_frame = pd.DataFrame(
        [
            {
                CODE_COLUMN: "THYAO",
                "Sektör": "Havacılık",
                "Güncel Fiyat": 100.0,
                "F/K Değerlemesi": 110.0,
                "PD/DD Finansal Model": 105.0,
                "EV/EBITDA Değerlemesi": 115.0,
                "DDM Değerlemesi": 100.0,
                "EFK Değerlemesi": 98.0,
                "NDK Değerlemesi": 102.0,
                "Graham Değerlemesi": 108.0,
            }
        ]
    )
    financials_frame = pd.DataFrame()
    dcf_frame = pd.DataFrame([{CODE_COLUMN: "THYAO", INA_VALUE_COLUMN: 120.0}])
    ratio_frame = pd.DataFrame()

    write_template_report(
        output_path=str(output_path),
        tickers=["THYAO.IS"],
        valuation_frame=valuation_frame,
        financials_frame=financials_frame,
        dcf_frame=dcf_frame,
        ratio_frame=ratio_frame,
        template_path=missing_template,
    )

    assert output_path.exists()
    workbook = load_workbook(output_path, read_only=True)
    assert SHEET_KALITE in workbook.sheetnames
    workbook.close()


def test_model_value_reason_marks_below_price_threshold() -> None:
    reason = _model_value_reason(raw_value=9.0, filtered_value=None, price=300.0)
    assert reason == "fiyata_gore_cok_dusuk"


def test_energy_equipment_sector_uses_dedicated_weight_profile() -> None:
    assert _weight_profile_for_sector("Enerji Ekipman & Taahhüt") == WEIGHT_PROFILE_ENERGY_EQUIPMENT


def test_weight_profile_mapping_is_generalized_across_sector_families() -> None:
    assert _weight_profile_for_sector("Banka") == WEIGHT_PROFILE_FINANCIAL
    assert _weight_profile_for_sector("Teknoloji") == WEIGHT_PROFILE_GROWTH
    assert _weight_profile_for_sector("Sanayi & Üretim") == WEIGHT_PROFILE_ENERGY_EQUIPMENT
    assert _weight_profile_for_sector("Gıda") == WEIGHT_PROFILE_DEFAULT


def test_sector_profile_config_contains_energy_equipment_mapping() -> None:
    config = load_sector_profile_config()
    energy_equipment = config["sector_groups"]["Enerji Ekipman/Taahhüt"]
    assert "Enerji Ekipman & Taahhüt" in energy_equipment


def test_valuation_profile_config_contains_growth_weights() -> None:
    config = load_valuation_profile_config()
    growth_weights = config["weights"]["Büyüme/Teknoloji"]
    assert growth_weights["D6"] > growth_weights["D1"]


def test_ratio_profile_config_contains_financial_rules() -> None:
    config = load_ratio_profile_config()
    financial_metrics = config["profiles"]["Finansal"]["metrics"]
    assert "roe" in financial_metrics
    assert financial_metrics["roe"]["kind"] == "band"


def test_ratio_frame_uses_sector_sensitive_profiles_and_0_100_scores() -> None:
    ratio_frame = pd.DataFrame(
        [
            {
                "Hisse": "MIATK",
                "Cari Oran": 2.0,
                "Likit Oran": 1.8,
                "Nakit Oran": 0.5,
                "Özsermaye Kârlılığı (ROE) (%) Yıllık": 18.0,
                "F/K": 26.0,
                "PD/DD": 5.5,
                "FAVÖK Büyüme (%) (Yıllık)": 30.0,
                "PEG Oranı": 1.2,
                "ROIC (%)": 13.0,
                "Aktif Devir Hızı": 0.7,
            },
            {
                "Hisse": "ULKER",
                "Cari Oran": 2.0,
                "Likit Oran": 1.8,
                "Nakit Oran": 0.5,
                "Özsermaye Kârlılığı (ROE) (%) Yıllık": 18.0,
                "F/K": 26.0,
                "PD/DD": 5.5,
                "FAVÖK Büyüme (%) (Yıllık)": 30.0,
                "PEG Oranı": 1.2,
                "ROIC (%)": 13.0,
                "Aktif Devir Hızı": 0.7,
            },
        ]
    )
    endeks_frame = pd.DataFrame(
        [
            {CODE_COLUMN: "MIATK", "Sektör": "Teknoloji"},
            {CODE_COLUMN: "ULKER", "Sektör": "Gıda"},
        ]
    )

    result = build_rasyo_frame(ratio_frame, endeks_frame)
    tech_row = result[result["Hisse"] == "MIATK"].iloc[0]
    food_row = result[result["Hisse"] == "ULKER"].iloc[0]

    assert tech_row[RASYO_PROFILE] == WEIGHT_PROFILE_GROWTH
    assert tech_row[RASYO_SCORE_100] is not None
    assert 0 <= tech_row[RASYO_SCORE_100] <= 100
    assert tech_row[RASYO_SCORE] is not None
    assert 0 <= tech_row[RASYO_SCORE] <= 1
    assert tech_row[RASYO_SCORE_100] > food_row[RASYO_SCORE_100]
    assert tech_row[RASYO_VALUATION_SCORE] is not None


def test_ratio_frame_penalizes_sparse_data_but_keeps_score() -> None:
    ratio_frame = pd.DataFrame(
        [
            {
                "Hisse": "AKBNK",
                "Özsermaye Kârlılığı (ROE) (%) Yıllık": 24.0,
                "F/K": 7.0,
                "PD/DD": 1.2,
            }
        ]
    )
    endeks_frame = pd.DataFrame([{CODE_COLUMN: "AKBNK", "Sektör": "Banka"}])

    result = build_rasyo_frame(ratio_frame, endeks_frame)
    row = result.iloc[0]

    assert row[RASYO_PROFILE] == WEIGHT_PROFILE_FINANCIAL
    assert row[RASYO_SCORE] is not None
    assert row[RASYO_COVERAGE] is not None
    assert row[RASYO_COVERAGE] < 1.0


def test_single_model_result_is_flagged_for_review() -> None:
    inputs = HisseInputs(
        code="TEST",
        sector_name="Enerji",
        index_name="XUTUM",
        market_key="tr",
        price=100.0,
        company_pe=10.0,
        company_pb=1.0,
        nis_sign="+",
        ratio_score=0.70,
        sector_pe=10.0,
        sector_pb=1.0,
        sector_ev_ebitda=8.0,
        eps=10.0,
        forward_eps=None,
        forward_pe=None,
        roe=0.15,
        beta=1.0,
        debt_ratio=0.20,
        net_income_growth=0.10,
        earnings_growth=0.10,
        revenue_growth=0.10,
        net_income=None,
        operating_income=None,
        equity=None,
        paid_in_capital=None,
        ebitda=None,
        free_cash_flow=None,
        operating_cash_flow=None,
        net_debt=0.0,
        asset_growth=None,
        dcf_value=None,
        analyst_target=None,
        analyst_count=None,
    )

    summary = valuation_summary({"D7": 120.0}, inputs, "Enerji Utility/Altyapı")

    assert summary.publishable is False
    assert summary.status == "İnceleme Gerekli"
    assert summary.note == "model_sayisi_yetersiz"


def test_d8_uses_bond_adjusted_multiple_instead_of_direct_bond_division() -> None:
    inputs = HisseInputs(
        code="TEST",
        sector_name="Sanayi",
        index_name="XUTUM",
        market_key="tr",
        price=100.0,
        company_pe=10.0,
        company_pb=1.0,
        nis_sign="+",
        ratio_score=0.70,
        sector_pe=10.0,
        sector_pb=1.0,
        sector_ev_ebitda=8.0,
        eps=10.0,
        forward_eps=None,
        forward_pe=None,
        roe=0.15,
        beta=1.0,
        debt_ratio=0.20,
        net_income_growth=0.10,
        earnings_growth=0.10,
        revenue_growth=0.10,
        net_income=100.0,
        operating_income=120.0,
        equity=500.0,
        paid_in_capital=10.0,
        ebitda=150.0,
        free_cash_flow=90.0,
        operating_cash_flow=110.0,
        net_debt=20.0,
        asset_growth=10.0,
        dcf_value=100.0,
        analyst_target=110.0,
        analyst_count=20.0,
    )
    macro_rates = MacroRates(
        market_key="tr",
        two_year_bond=0.40,
        market_premium=0.08,
        terminal_growth=0.03,
        cost_of_debt=0.30,
        tax_rate=0.25,
        risk_free_rate=0.25,
    )

    values = _raw_valuation_points(inputs, macro_rates)

    assert math.isclose(values["D1"], 100.0)
    assert math.isclose(values["D8"], 50.0)


def test_valuation_summary_keeps_high_but_valid_models_in_average() -> None:
    inputs = HisseInputs(
        code="ASELS",
        sector_name="Savunma",
        index_name="XUTUM",
        market_key="tr",
        price=420.25,
        company_pe=64.06,
        company_pb=7.65,
        nis_sign="+",
        ratio_score=0.5942,
        sector_pe=55.66,
        sector_pb=5.47,
        sector_ev_ebitda=30.03,
        eps=7.55,
        forward_eps=11.66,
        forward_pe=64.06,
        roe=0.18,
        beta=1.0,
        debt_ratio=0.15,
        net_income_growth=0.10,
        earnings_growth=0.10,
        revenue_growth=0.10,
        net_income=100.0,
        operating_income=120.0,
        equity=549.19,
        paid_in_capital=10.0,
        ebitda=150.0,
        free_cash_flow=90.0,
        operating_cash_flow=110.0,
        net_debt=28.13,
        asset_growth=17.0,
        dcf_value=97.24,
        analyst_target=351.23,
        analyst_count=10.0,
    )

    summary = valuation_summary(
        {
            "D1": 420.23,
            "D2": 90.73,
            "D3": 65.68,
            "D4": 420.13,
            "D5": 54.92,
            "D6": 747.01,
            "D7": 422.37,
            "D8": 210.17,
            "D9": 22.84,
            "D10": None,
            "D11": 97.24,
            "D12": None,
        },
        inputs,
        WEIGHT_PROFILE_ENERGY_EQUIPMENT,
    )

    assert summary.fair_value is not None
    assert summary.fair_value > 250.0


def test_valuation_summary_marks_extreme_price_and_consensus_mismatch_for_review() -> None:
    inputs = HisseInputs(
        code="TEST",
        sector_name="Sanayi",
        index_name="XUTUM",
        market_key="tr",
        price=420.0,
        company_pe=12.0,
        company_pb=1.5,
        nis_sign="+",
        ratio_score=0.70,
        sector_pe=12.0,
        sector_pb=1.5,
        sector_ev_ebitda=8.0,
        eps=10.0,
        forward_eps=10.5,
        forward_pe=12.0,
        roe=0.18,
        beta=1.0,
        debt_ratio=0.15,
        net_income_growth=0.10,
        earnings_growth=0.10,
        revenue_growth=0.10,
        net_income=100.0,
        operating_income=120.0,
        equity=500.0,
        paid_in_capital=10.0,
        ebitda=150.0,
        free_cash_flow=90.0,
        operating_cash_flow=110.0,
        net_debt=20.0,
        asset_growth=10.0,
        dcf_value=95.0,
        analyst_target=350.0,
        analyst_count=12.0,
    )

    summary = valuation_summary(
        {"D2": 90.0, "D3": 92.0, "D5": 88.0, "D8": 100.0, "D11": 95.0},
        inputs,
        WEIGHT_PROFILE_DEFAULT,
    )

    assert summary.publishable is False
    assert summary.status == "İnceleme Gerekli"
    assert summary.note == "piyasa_konsensus_uyumsuzlugu"


def test_valuation_summary_stays_publishable_when_price_and_consensus_are_reasonably_aligned() -> None:
    inputs = HisseInputs(
        code="TEST",
        sector_name="Sanayi",
        index_name="XUTUM",
        market_key="tr",
        price=420.0,
        company_pe=12.0,
        company_pb=1.5,
        nis_sign="+",
        ratio_score=0.70,
        sector_pe=12.0,
        sector_pb=1.5,
        sector_ev_ebitda=8.0,
        eps=10.0,
        forward_eps=10.5,
        forward_pe=12.0,
        roe=0.18,
        beta=1.0,
        debt_ratio=0.15,
        net_income_growth=0.10,
        earnings_growth=0.10,
        revenue_growth=0.10,
        net_income=100.0,
        operating_income=120.0,
        equity=500.0,
        paid_in_capital=10.0,
        ebitda=150.0,
        free_cash_flow=90.0,
        operating_cash_flow=110.0,
        net_debt=20.0,
        asset_growth=10.0,
        dcf_value=410.0,
        analyst_target=390.0,
        analyst_count=12.0,
    )

    summary = valuation_summary(
        {"D1": 420.0, "D4": 400.0, "D7": 430.0, "D8": 390.0, "D11": 410.0},
        inputs,
        WEIGHT_PROFILE_DEFAULT,
    )

    assert summary.publishable is True
    assert summary.fair_value is not None


def test_quality_inputs_improve_confidence_when_ratio_and_nis_are_strong() -> None:
    strong_inputs = HisseInputs(
        code="TEST",
        sector_name="Sanayi",
        index_name="XUTUM",
        market_key="tr",
        price=100.0,
        company_pe=10.0,
        company_pb=1.0,
        nis_sign="+",
        ratio_score=0.70,
        sector_pe=10.0,
        sector_pb=1.0,
        sector_ev_ebitda=8.0,
        eps=10.0,
        forward_eps=None,
        forward_pe=None,
        roe=0.15,
        beta=1.0,
        debt_ratio=0.20,
        net_income_growth=0.10,
        earnings_growth=0.10,
        revenue_growth=0.10,
        net_income=100.0,
        operating_income=120.0,
        equity=500.0,
        paid_in_capital=10.0,
        ebitda=150.0,
        free_cash_flow=90.0,
        operating_cash_flow=110.0,
        net_debt=20.0,
        asset_growth=10.0,
        dcf_value=100.0,
        analyst_target=110.0,
        analyst_count=20.0,
        revenue=1000.0,
        gross_margin=40.0,
        total_assets=2000.0,
        total_debt=600.0,
    )
    weak_inputs = HisseInputs(
        **{
            **strong_inputs.__dict__,
            "nis_sign": "-",
            "ratio_score": 0.35,
            "gross_margin": 12.0,
            "total_debt": 1500.0,
        }
    )

    strong_summary = valuation_summary(
        {"D1": 100.0, "D7": 105.0, "D11": 98.0},
        strong_inputs,
        WEIGHT_PROFILE_DEFAULT,
    )
    weak_summary = valuation_summary(
        {"D1": 100.0, "D7": 105.0, "D11": 98.0},
        weak_inputs,
        WEIGHT_PROFILE_DEFAULT,
    )

    assert strong_summary.confidence is not None
    assert weak_summary.confidence is not None
    assert strong_summary.confidence > weak_summary.confidence


def test_hisse_output_row_includes_fundamental_quality_score() -> None:
    inputs = HisseInputs(
        code="TEST",
        sector_name="Sanayi",
        index_name="XUTUM",
        market_key="tr",
        price=100.0,
        company_pe=10.0,
        company_pb=1.0,
        nis_sign="+",
        ratio_score=0.70,
        sector_pe=10.0,
        sector_pb=1.0,
        sector_ev_ebitda=8.0,
        eps=10.0,
        forward_eps=None,
        forward_pe=None,
        roe=0.15,
        beta=1.0,
        debt_ratio=0.20,
        net_income_growth=0.10,
        earnings_growth=0.10,
        revenue_growth=0.10,
        net_income=100.0,
        operating_income=120.0,
        equity=500.0,
        paid_in_capital=10.0,
        ebitda=150.0,
        free_cash_flow=90.0,
        operating_cash_flow=110.0,
        net_debt=20.0,
        asset_growth=10.0,
        dcf_value=100.0,
        analyst_target=110.0,
        analyst_count=20.0,
        revenue=1000.0,
        gross_margin=40.0,
        total_assets=2000.0,
        total_debt=600.0,
    )

    row = _hisse_output_row(inputs, {"D1": 100.0, "D7": 105.0, "D11": 98.0})

    assert row[HISSE_FUNDAMENTAL_QUALITY] is not None
    assert row[HISSE_FUNDAMENTAL_QUALITY] > 0.60


def test_us_market_profile_has_no_is_suffix() -> None:
    assert US_PROFILE.ticker_suffix == ""


def test_us_valuation_overlay_changes_growth_weights() -> None:
    config = load_valuation_profile_config()
    us_weights = resolve_valuation_weight_map(config, "Büyüme/Teknoloji", "us")
    tr_weights = resolve_valuation_weight_map(config, "Büyüme/Teknoloji", "tr")
    assert us_weights["D11"] > tr_weights["D11"]
    assert us_weights["D6"] < tr_weights["D6"]


def test_us_ratio_overlay_changes_growth_thresholds() -> None:
    config = load_ratio_profile_config()
    us_profile = resolve_ratio_profile(config, "Büyüme/Teknoloji", "us")
    tr_profile = resolve_ratio_profile(config, "Büyüme/Teknoloji", "tr")
    assert us_profile["metrics"]["pe"]["target_max"] > tr_profile["metrics"]["pe"]["target_max"]
